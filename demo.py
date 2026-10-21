#!/usr/bin/env python3
"""
End-to-end demo on FICTIONAL data -- runs without Claude, without any account,
in about 20 seconds.

    pip install -r requirements.txt
    python demo.py

What it does, in the order the studio would:
  1. generate the inputs: a generic legend template, 3 "old project" legends, spec-sheet PDFs
  2. build the fixture bank from the 3 old legends (duplicates across projects are merged)
  3. build the selection page (examples/out/picker.html)
  4. export 4 fixtures chosen from the bank into a NEW project legend
  5. add 3 fixtures from spec-sheet PDFs  <- the LLM step: Claude's recorded decisions are
     replayed from examples/data/claude_decisions.json; everything else is plain Python
  6. check each new fixture against the bank ("already there" / "new")
  7. the reviewer returns the file -> new fixtures join the bank, but only once nothing is
     still orange (unverified)

Outputs land in examples/out/ (git-ignored). Open the .xlsx files in Excel to see the result.
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "examples" / "out"
PY = sys.executable
sys.path.insert(0, str(ROOT / "skill" / "scripts"))
sys.path.insert(0, str(ROOT / "bank"))

for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")

SUPPLIERS = {  # fictional local importers
    "VELTRANO": "ספק א׳ (דמו)", "CORVELLA": "ספק א׳ (דמו)", "PRAVENA": "ספק ב׳ (דמו)",
    "PRAVENNA": "ספק ב׳ (דמו)", "OSKARI LUX": "ספק ג׳ (דמו)", "HALVARD LUMEN": "ספק ב׳ (דמו)",
}
NEW_PROJECT = "דירת דמו"
NEW_LEGEND = OUT / f"מקרא - {NEW_PROJECT}.xlsx"


def step(n, title):
    print(f"\n{'=' * 78}\n{n}. {title}\n{'=' * 78}")


def run(*args, expect_fail=False, quiet=False):
    shown = " ".join(str(a).replace(str(ROOT) + "/", "").replace(str(ROOT) + "\\", "") for a in args)
    print(f"$ python {shown}")
    r = subprocess.run([PY, *map(str, args)], cwd=ROOT, capture_output=True, text=True,
                       encoding="utf-8", errors="replace", env={"PYTHONIOENCODING": "utf-8",
                                                                **__import__("os").environ})
    out = (r.stdout + r.stderr).strip()
    if out and not quiet:
        print("  " + out.replace("\n", "\n  ").replace(str(ROOT) + "/", "").replace(str(ROOT) + "\\", ""))
    if (r.returncode != 0) != expect_fail:
        sys.exit(f"\n!! unexpected exit code {r.returncode} -- demo stopped")
    return out


def bank_row_of(model):
    import migrate_mikra_to_bank as M
    wb = M.open_workbook(OUT / "bank.xlsx", sheet=M.BANK_SHEET_NAME)
    ws = wb[M.BANK_SHEET_NAME]
    hr = M.find_bank_header_row(ws)
    for r in range(hr + 1, ws.max_row + 1):
        if f"מדגם {model}\n" in str(ws.cell(row=r, column=M.col_index("תיאור")).value or "") + "\n":
            return r
    sys.exit(f"model {model!r} not found in the bank")


def bank_size():
    import migrate_mikra_to_bank as M
    wb = M.open_workbook(OUT / "bank.xlsx", sheet=M.BANK_SHEET_NAME)
    ws = wb[M.BANK_SHEET_NAME]
    hr = M.find_bank_header_row(ws)
    return sum(1 for r in range(hr + 1, ws.max_row + 1) if ws.cell(row=r, column=M.col_index("יצרן")).value)


def reviewer_fixes(path, row):
    """What the person reviewing the file does: confirms the orange values and removes the colour."""
    import openpyxl
    from openpyxl.styles import PatternFill
    import insert_fixture as IF
    wb = openpyxl.load_workbook(path, rich_text=True)
    ws = wb[IF.SHEET_NAME]
    hm = IF.build_header_map(ws)
    c = ws.cell(row=row, column=hm["גוון"])
    c.value, c.fill, c.comment = "אנתרציט", PatternFill(fill_type=None), None   # finish now chosen
    wb.save(path)


def main():
    if OUT.exists():
        shutil.rmtree(OUT)

    step(1, "Generate fictional inputs (template, 3 old project legends, spec-sheet PDFs)")
    run("examples/make_sample_data.py", quiet=True)
    for p in sorted((OUT / "old_projects").glob("*.xlsx")):
        print(f"  old project: {p.name}")

    step(2, "Build the fixture bank from the old legends")
    run("bank/build_bank_template.py", OUT / "bank.xlsx")
    sup = [a for k, v in SUPPLIERS.items() for a in ("--supplier", f"{k}={v}")]
    for p in sorted((OUT / "old_projects").glob("*.xlsx")):
        run("bank/migrate_mikra_to_bank.py", p, OUT / "bank.xlsx", *sup)
    old_rows = sum(len(p["rows"]) for p in json.loads(
        (ROOT / "examples/data/fixtures.json").read_text(encoding="utf8"))["old_projects"])
    print(f"\n  -> {old_rows} rows in 3 old legends became {bank_size()} fixtures in the bank")

    step(3, "Build the selection page from the bank")
    run("pages/build_pages.py", "--bank", OUT / "bank.xlsx", "--outdir", OUT, "--demo-note")

    step(4, f"Export 4 fixtures from the bank into a new project: '{NEW_PROJECT}', 3000K")
    rows = [bank_row_of(m) for m in ("Kito 24°", "Strada 20", "Fjell Up/Down", "Terra 60")]
    run("bank/export_selection.py", OUT / "bank.xlsx", "--rows", ",".join(map(str, rows)),
        "--name", NEW_PROJECT, "--color", "3000K",
        "--skeleton", OUT / "legend_template.xlsx", "--out", NEW_LEGEND)

    step(5, "Add fixtures from spec-sheet PDFs (the LLM step, replayed)")
    import insert_fixture as IF
    import match_fixture as MF
    decisions = json.loads((ROOT / "examples/data/claude_decisions.json").read_text(encoding="utf8"))
    added_rows = []
    for d in decisions["spec_sheets"]:
        pdf = ROOT / d["pdf"]
        extract = OUT / "pdf_extract" / pdf.stem
        text = run("skill/scripts/read_pdf.py", pdf, "--out", extract, quiet=True)
        print(f"  read {len(text.split())} words of text from {pdf.name}")
        for q in d["questions_asked_in_one_batch"]:
            print(f"  Claude asked: {q['question']}  ->  {q['answer']}")
        row = IF.insert_fixture(str(NEW_LEGEND), d["category"], d["data"],
                                image_path=str(extract / d["image"]),
                                uncertain_fields=d["uncertain_fields"],
                                uncertain_snippets=d["uncertain_snippets"])
        step_label = d["data"]["תיאור"].split("\n")[0]

        # 6. bank check, right after each insert (SKILL.md step 3.5)
        verdict = MF.find_match(d["data"]["יצרן"], d["data"]["תיאור"], OUT / "bank.xlsx")
        print(f"  bank check for '{step_label}': {verdict['verdict']} (score {verdict['score']})")
        if verdict["verdict"] in ("new", "maybe"):
            extra = ["--bank-description", verdict["bank_description"]] if verdict["verdict"] == "maybe" else []
            run("bank/mark_bank_status.py", NEW_LEGEND, "--row", row, "--status", verdict["verdict"], *extra)
        added_rows.append((row, verdict["verdict"]))

    snapshot = OUT / f"מקרא - {NEW_PROJECT} (לפני בדיקה).xlsx"
    shutil.copy(NEW_LEGEND, snapshot)
    print(f"\n  snapshot of the file as it goes to the reviewer: {snapshot.name}")

    step(7, "The reviewer sends the file back")
    before = bank_size()
    print("  a) returned WITHOUT checking the orange cells:")
    run("bank/process_new_fixtures.py", NEW_LEGEND, OUT / "bank.xlsx", "--no-rebuild")
    assert bank_size() == before, "an unverified fixture reached the bank"
    new_row = next(r for r, v in added_rows if v == "new")
    print("\n  b) reviewer confirms the values and removes the orange colour, returns it again:")
    reviewer_fixes(NEW_LEGEND, new_row)
    run("bank/process_new_fixtures.py", NEW_LEGEND, OUT / "bank.xlsx")
    print("\n  c) the same file is returned twice by mistake:")
    run("bank/process_new_fixtures.py", NEW_LEGEND, OUT / "bank.xlsx", "--no-rebuild")
    after = bank_size()
    assert after == before + 1, (before, after)

    # the rebuild in 7b refreshed the page; add the demo banner back for viewing
    run("pages/build_pages.py", "--bank", OUT / "bank.xlsx", "--outdir", OUT, "--demo-note", quiet=True)

    print(f"\n{'=' * 78}\nDone. Bank: {before} -> {after} fixtures.")
    print(f"Open in Excel:  examples/out/{NEW_LEGEND.name}  and  examples/out/bank.xlsx")
    print("Open in a browser:  examples/out/picker.html")


if __name__ == "__main__":
    main()
