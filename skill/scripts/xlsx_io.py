"""
Open and save legend / bank workbooks so that failures are LOUD and SAFE.

- open_workbook(): a missing, corrupt or non-Excel file raises WorkbookError with a
  plain-Hebrew message, instead of a raw zipfile/openpyxl traceback.
- save_workbook(): writes to a temporary file in the same folder first, then
  swaps it into place. If the save fails half-way (disk full, crash), the
  original file is untouched -- never a half-written .xlsx. If the target is
  open in Excel on Windows (Excel locks the file), the swap fails with a clear
  "close the file and try again" message, and the original is untouched.
"""
import os
import tempfile
import zipfile
from pathlib import Path

import openpyxl


class WorkbookError(RuntimeError):
    """A user-facing problem with a workbook file (message is meant to be shown as-is)."""


def open_workbook(path, sheet=None, **kwargs):
    p = Path(path)
    if not p.exists():
        raise WorkbookError(f"הקובץ לא נמצא: {p}")
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise WorkbookError(f"זה לא קובץ Excel (.xlsx): {p.name}")
    # rich_text=True by default: without it openpyxl silently flattens a cell where
    # only one phrase is marked orange ("2700K" inside a description) into plain
    # text on the next save -- so any later script that merely adds a comment to
    # the file would erase the "check this phrase" marking.
    if not kwargs.get("read_only"):
        kwargs.setdefault("rich_text", True)
    try:
        wb = openpyxl.load_workbook(p, **kwargs)
    except PermissionError:
        raise WorkbookError(f"אין גישה לקובץ {p.name} — ייתכן שהוא פתוח באקסל או מסונכרן כרגע. "
                            "סגרו אותו ונסו שוב.") from None
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
        raise WorkbookError(f"הקובץ {p.name} פגום או לא נשמר עד הסוף, ולא ניתן לפתוח אותו ({e.__class__.__name__}). "
                            "לא בוצע שום שינוי. אפשר לשחזר מגיבוי או לפתוח ולשמור מחדש באקסל.") from None
    if sheet is not None and sheet not in wb.sheetnames:
        raise WorkbookError(f"בקובץ {p.name} אין גיליון בשם '{sheet}' (יש: {', '.join(wb.sheetnames)}). "
                            "ייתכן שזה לא הקובץ הנכון.")
    return wb


def save_workbook(wb, path):
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".~saving-", suffix=".xlsx", dir=str(p.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        try:
            os.replace(tmp, p)
        except PermissionError:
            raise WorkbookError(f"לא ניתן לשמור את {p.name} — הקובץ פתוח כנראה באקסל. "
                                "סגרו אותו והריצו שוב. הקובץ המקורי לא השתנה.") from None
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
    return p


def replace_file(src, dst):
    """Copy a finished file over dst atomically (same temp-then-swap rule as save_workbook)."""
    import shutil
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".~saving-", suffix=dst.suffix, dir=str(dst.parent))
    os.close(fd)
    try:
        shutil.copyfile(src, tmp)
        try:
            os.replace(tmp, dst)
        except PermissionError:
            raise WorkbookError(f"לא ניתן לשמור את {dst.name} — הקובץ פתוח כנראה באקסל. "
                                "סגרו אותו והריצו שוב. הקובץ המקורי לא השתנה.") from None
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
    return dst
