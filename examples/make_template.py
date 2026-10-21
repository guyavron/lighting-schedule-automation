#!/usr/bin/env python3
"""
Build a generic, empty lighting legend template (מקרא תאורה) from code.

It has the same STRUCTURE the scripts expect from a real studio template --
sheet "אומדן", a project header (rows 1-7), the column-header row on row 8,
three category blocks (ceiling / wall / floor) with an instructional example
row, and a summary row with a SUM formula -- but generic styling, no logo and
no studio branding. Built in code rather than shipped as a binary .xlsx so you
can read exactly what's in it.

    python examples/make_template.py examples/out/legend_template.xlsx
"""
import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET = "אומדן"
HEADER_ROW = 8
COLUMNS = [  # header text, width
    ("סעיף", 6), ("סימון", 14), ("תיאור", 42), ("יצרן", 16), ("מתח/\nפרוטוקול עמעום", 16),
    ("גוון", 14), ("תמונה", 22), ("כמות", 8), ('סה"כ יחידות', 9), ("אורך פרופיל (מטרים)", 11),
    ("מחיר למטר/ יחידה", 11), ("מחיר פרופילים", 11), ('סה"כ', 12),
]
# Category header text is what insert_fixture.py searches for -- must match CATEGORY_TEXT there.
CATEGORIES = [("מנורות הזנה מהתקרה", "D9C4E8"), ("מנורות הזנה מהקיר", "C4D9E8"),
              ("מנורות הזנה מהרצפה", "C9E8C4")]
SKELETON = {  # the bracketed example row; insert_fixture.py strips it on the first save
    "סימון": "[קיצור סימון לתוכנית, למשל SP / TR-XXX]",
    "תיאור": "[סוג הגוף] מדגם [שם הדגם]\nLED [הספק]W [שטף אור]LM [גוון]K CRI[ערך]\n"
             "זוית פיזור [מעלות] מעלות\nדרייברים למיקום באזור פריק ונגיש",
    "יצרן": "[שם היצרן באנגלית]",
    "מתח/\nפרוטוקול עמעום": "[מתח]\n[פרוטוקול עמעום, למשל DIM DALI]",
    "גוון": "[גוון גימור הגוף]",
    "תמונה": "[תמונת המוצר מגיליון היצרן]",
}

THIN = Side(style="thin", color="9AA3AB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
WRAP_R = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
HEAD_FILL = PatternFill("solid", start_color="E9ECEF", end_color="E9ECEF")


def build(out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.sheet_view.rightToLeft = True
    ncol = len(COLUMNS)

    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    # project header -- set_project_header.py writes C2 / C3 / C5
    ws["A1"] = "מקרא גופי תאורה"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["C2"] = "שם הפרויקט:  [שם פרויקט]"
    ws["C3"] = "תאריך:[תאריך]"
    ws["C5"] = "גוון תאורה בפרוייקט: [2700K/3000K]"
    ws.merge_cells("C5:D6")
    for c in ("C2", "C3", "C5"):
        ws[c].font = BOLD
        ws[c].alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)

    for i, (text, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=i, value=text)
        cell.font, cell.border, cell.alignment, cell.fill = BOLD, BOX, CENTER, HEAD_FILL
    ws.row_dimensions[HEADER_ROW].height = 30

    row = HEADER_ROW + 1
    for idx, (title, colour) in enumerate(CATEGORIES):
        ws.cell(row=row, column=1, value=title).font = BOLD
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        bar = PatternFill("solid", start_color=colour, end_color=colour)
        for c in range(1, ncol + 1):
            ws.cell(row=row, column=c).fill = bar
            ws.cell(row=row, column=c).border = BOX
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
        ws.row_dimensions[row].height = 20
        row += 1
        body_rows = 2
        for k in range(body_rows):
            for c in range(1, ncol + 1):
                cell = ws.cell(row=row, column=c)
                cell.border, cell.font, cell.alignment = BOX, FONT, WRAP_R
            if idx == 0 and k == 0:
                ws.cell(row=row, column=1, value=1)
                for c, (text, _) in enumerate(COLUMNS, start=1):
                    if text in SKELETON:
                        ws.cell(row=row, column=c, value=SKELETON[text])
            ws.row_dimensions[row].height = 30.75
            row += 1

    summary = row
    ws.cell(row=summary, column=3, value='סה"כ כולל נורות וציוד:').font = BOLD
    ws.cell(row=summary, column=ncol, value=f"=SUM(M{HEADER_ROW + 2}:M{summary - 1})").font = BOLD
    ws.freeze_panes = f"A{HEADER_ROW + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True   # print all columns on one page width
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out_xlsx", nargs="?", default="examples/out/legend_template.xlsx")
    args = ap.parse_args()
    print(f"Template written: {build(args.out_xlsx)}")


if __name__ == "__main__":
    sys.exit(main())
