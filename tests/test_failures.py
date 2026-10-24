"""
Failure-mode tests: what happens when the input is wrong, not only when it's right.

The rule they enforce: a visible error (or an orange cell) is acceptable; a silent
wrong result -- an invented value, an empty cell that looks finished, a half-written
file -- is not.

    python -m unittest discover -s tests -v
"""
import contextlib
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skill" / "scripts"))
sys.path.insert(0, str(ROOT / "bank"))
sys.path.insert(0, str(ROOT / "examples"))

import openpyxl                              # noqa: E402
from openpyxl.cell.rich_text import CellRichText  # noqa: E402

import insert_fixture as IF                  # noqa: E402
import xlsx_io                               # noqa: E402
import make_template                         # noqa: E402
import draw_products                         # noqa: E402
import migrate_mikra_to_bank as M            # noqa: E402
import match_fixture as MF                   # noqa: E402
import build_bank_template as BBT            # noqa: E402
from bank_guard import bank_guard            # noqa: E402

SPECS = ROOT / "examples" / "spec-sheets"
PY = sys.executable
ENV = {**os.environ, "PYTHONIOENCODING": "utf-8"}

FX = {"תיאור": "ספוט שקוע מדגם Testa 10\nLED 8W 700LM 3000K CRI90\nדרייברים למיקום באזור פריק ונגיש",
      "יצרן": "VELTRANO", "מתח_עמעום": "350mA\nDIM DALI", "גוון": "לבן"}


def setUpModule():
    if not (SPECS / "broken" / "truncated.pdf").exists():
        subprocess.run([PY, str(ROOT / "examples" / "make_sample_data.py")], check=True,
                       capture_output=True, env=ENV)


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def run(*args):
    r = subprocess.run([PY, *map(str, args)], cwd=ROOT, capture_output=True, text=True,
                       encoding="utf-8", errors="replace", env=ENV)
    return r.returncode, r.stdout + r.stderr


def orange(cell):
    return cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb).upper().endswith("FFC000")


class Tmp(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        self.legend = make_template.build(self.d / "legend.xlsx")
        self.photo = draw_products.draw("downlight", "לבן", self.d / "p.jpg")

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def sheet(self, path=None):
        return openpyxl.load_workbook(path or self.legend, rich_text=True)[IF.SHEET_NAME]

    def cell(self, ws, row, field):
        return ws.cell(row=row, column=IF.build_header_map(ws)[field])


# ------------------------------------------------------------------ reading PDFs
class TestReadPdf(Tmp):
    def test_good_pdf_returns_text(self):
        code, out = run("skill/scripts/read_pdf.py", SPECS / "corvella_track-one-15.pdf", "--out", self.d / "x")
        self.assertEqual(code, 0)
        self.assertIn("Luminous flux", out)

    def test_scanned_pdf_is_a_loud_error_not_empty_success(self):
        code, out = run("skill/scripts/read_pdf.py", SPECS / "broken" / "scanned_no-text-layer.pdf", "--out", self.d / "x")
        self.assertEqual(code, 5)
        self.assertIn("no text layer", out)
        self.assertTrue(list((self.d / "x").glob("page*.jpg")), "page images should still be rendered for a human")

    def test_truncated_pdf(self):
        code, out = run("skill/scripts/read_pdf.py", SPECS / "broken" / "truncated.pdf", "--out", self.d / "x")
        self.assertEqual(code, 4)
        self.assertIn("damaged", out)
        self.assertNotIn("Traceback", out)

    def test_not_a_pdf(self):
        code, out = run("skill/scripts/read_pdf.py", SPECS / "broken" / "not-really-a.pdf", "--out", self.d / "x")
        self.assertEqual(code, 3)
        self.assertNotIn("Traceback", out)

    def test_missing_file(self):
        code, _ = run("skill/scripts/read_pdf.py", self.d / "nope.pdf")
        self.assertEqual(code, 2)


# ------------------------------------------------------------- writing a legend row
class TestInsertFixture(Tmp):
    def test_happy_path_fields_and_image(self):
        row = IF.insert_fixture(str(self.legend), "ceiling", FX, image_path=str(self.photo))
        ws = self.sheet()
        self.assertEqual(self.cell(ws, row, "יצרן").value, "VELTRANO")
        self.assertEqual(len(ws._images), 1)
        self.assertFalse(any(orange(self.cell(ws, row, f)) for f in ("יצרן", "גוון", "תיאור")))

    def test_empty_field_is_flagged_not_left_looking_finished(self):
        row = IF.insert_fixture(str(self.legend), "ceiling", {**FX, "גוון": ""}, image_path=str(self.photo))
        c = self.cell(self.sheet(), row, "גוון")
        self.assertTrue(orange(c))
        self.assertIn("חסר", c.comment.text)

    def test_field_not_given_at_all_is_flagged(self):
        data = {k: v for k, v in FX.items() if k != "מתח_עמעום"}
        row = IF.insert_fixture(str(self.legend), "wall", data, image_path=str(self.photo))
        self.assertTrue(orange(self.cell(self.sheet(), row, "מתח_עמעום")))

    def test_missing_image_is_flagged(self):
        row = IF.insert_fixture(str(self.legend), "floor", FX)
        self.assertTrue(orange(self.cell(self.sheet(), row, "תמונה")))

    def test_snippet_not_found_flags_whole_cell(self):
        row = IF.insert_fixture(str(self.legend), "ceiling", FX, image_path=str(self.photo),
                                uncertain_snippets={"תיאור": "זוית פיזור 72 מעלות"})
        self.assertTrue(orange(self.cell(self.sheet(), row, "תיאור")))

    def test_unknown_category_changes_nothing(self):
        before = sha(self.legend)
        with self.assertRaises(ValueError):
            IF.insert_fixture(str(self.legend), "roof", FX)
        self.assertEqual(sha(self.legend), before)

    def test_unsupported_field_changes_nothing(self):
        before = sha(self.legend)
        with self.assertRaises(ValueError):
            IF.insert_fixture(str(self.legend), "ceiling", {**FX, "מחיר": "100"})
        self.assertEqual(sha(self.legend), before)

    def test_corrupt_xlsx_is_a_clear_error(self):
        bad = self.d / "bad.xlsx"
        bad.write_bytes(self.legend.read_bytes()[:500])
        before = sha(bad)
        code, out = run("skill/scripts/insert_fixture.py", bad, "--category", "ceiling", "--data", json.dumps(FX))
        self.assertNotEqual(code, 0)
        self.assertIn("פגום", out)
        self.assertNotIn("Traceback", out)
        self.assertEqual(sha(bad), before)

    def test_wrong_workbook_is_a_clear_error(self):
        other = self.d / "other.xlsx"
        openpyxl.Workbook().save(other)
        with self.assertRaises(xlsx_io.WorkbookError):
            IF.insert_fixture(str(other), "ceiling", FX)

    def test_file_open_in_excel_leaves_original_intact(self):
        """Excel on Windows locks an open file; replacing it raises PermissionError."""
        before = sha(self.legend)
        with mock.patch("xlsx_io.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(xlsx_io.WorkbookError) as e:
                IF.insert_fixture(str(self.legend), "ceiling", FX, image_path=str(self.photo))
        self.assertIn("פתוח", str(e.exception))
        self.assertEqual(sha(self.legend), before)
        self.assertEqual(list(self.d.glob(".~saving-*")), [], "temp file left behind")

    def test_crash_during_save_leaves_original_intact(self):
        before = sha(self.legend)
        with mock.patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                IF.insert_fixture(str(self.legend), "ceiling", FX)
        self.assertEqual(sha(self.legend), before)

    def test_new_row_does_not_inherit_orange_from_row_above(self):
        """Found in testing: a full block copied the style of the flagged row above."""
        rows = [IF.insert_fixture(str(self.legend), "ceiling", {**FX, "גוון": g}, image_path=str(self.photo))
                for g in ("", "שחור", "לבן")]
        ws = self.sheet()
        flags = [orange(self.cell(ws, r, "גוון")) for r in rows]
        self.assertEqual(flags, [True, False, False])

    def test_orange_phrase_survives_later_scripts(self):
        """Found in testing: mark_bank_status re-saved the file without rich text."""
        row = IF.insert_fixture(str(self.legend), "ceiling", FX, image_path=str(self.photo),
                                uncertain_snippets={"תיאור": "3000K"})
        code, _ = run("bank/mark_bank_status.py", self.legend, "--row", row, "--status", "new")
        self.assertEqual(code, 0)
        self.assertIsInstance(self.cell(self.sheet(), row, "תיאור").value, CellRichText)


# --------------------------------------------------------------------- the bank
class BankBase(Tmp):
    def make_bank(self, rows):
        """rows: list of (category_en, data) -> builds an old legend and migrates it into a new bank."""
        old = self.d / "old.xlsx"
        shutil.copy(self.legend, old)
        for cat, data in rows:
            IF.insert_fixture(str(old), cat, data, image_path=str(self.photo))
        bank = self.d / "bank.xlsx"
        BBT.build(bank)
        M.migrate(str(old), str(bank), str(bank), {})
        return bank


class TestBank(BankBase):
    def test_variants_with_special_characters_are_not_merged(self):
        a = {**FX, "תיאור": "שקוע תקרה מדגם Arca Mini Ø55\nLED 7W 640LM"}
        b = {**FX, "תיאור": "שקוע תקרה מדגם Arca Mini Ø55 IP65\nLED 7W 610LM"}
        bank = self.make_bank([("ceiling", a), ("ceiling", b)])
        ws = openpyxl.load_workbook(bank)[M.BANK_SHEET_NAME]
        hr = M.find_bank_header_row(ws)
        n = sum(1 for r in range(hr + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value)
        self.assertEqual(n, 2)

    def test_line_breaks_survive_the_bank(self):
        bank = self.make_bank([("ceiling", FX)])
        ws = openpyxl.load_workbook(bank)[M.BANK_SHEET_NAME]
        hr = M.find_bank_header_row(ws)
        self.assertEqual(ws.cell(row=hr + 1, column=M.col_index("תיאור")).value, FX["תיאור"])
        self.assertEqual(ws.cell(row=hr + 1, column=M.col_index("מתח_עמעום")).value, FX["מתח_עמעום"])

    def test_match_different_product_same_brand_is_not_strong(self):
        """Found in testing: 'Beam 30' scored 0.82 ('already in the bank') against 'Terra 60'."""
        terra = {**FX, "יצרן": "HALVARD LUMEN", "תיאור": "שקוע רצפה מדגם Terra 60\nLED 5W 410LM 3000K CRI90 IP67\nזוית פיזור 40 מעלות\nדרייברים למיקום באזור פריק ונגיש"}
        bank = self.make_bank([("floor", terra)])
        beam = "פנס הצפה לחזית מדגם Beam 30\nLED 20W 2100LM 3000K CRI80 IP66\nזוית פיזור 30 מעלות\nדרייברים למיקום באזור פריק ונגיש"
        self.assertNotEqual(MF.find_match("HALVARD LUMEN", beam, bank)["verdict"], "strong")
        self.assertEqual(MF.find_match("HALVARD LUMEN", terra["תיאור"], bank)["verdict"], "strong")

    def test_match_unknown_brand_is_new(self):
        bank = self.make_bank([("ceiling", FX)])
        self.assertEqual(MF.find_match("TOBEK", "גוף תאורה לינארי לקיר מדגם Riva 120\nLED 14W", bank)["verdict"], "new")

    def test_concurrent_run_is_blocked_and_bank_untouched(self):
        bank = self.make_bank([("ceiling", FX)])
        before = sha(bank)
        with bank_guard(bank, "first run"):
            code, out = run("bank/process_new_fixtures.py", self.legend, bank, "--no-rebuild")
        self.assertNotEqual(code, 0)
        self.assertIn("תפוס", out)
        self.assertEqual(sha(bank), before)

    def test_stale_lock_is_ignored(self):
        bank = self.make_bank([("ceiling", FX)])
        state = bank.parent / ".bank_state.json"
        state.write_text(json.dumps({"lock": {"who": "x", "action": "crashed", "ts": time.time() - 3600}}), encoding="utf8")
        with bank_guard(bank, "after a crash"):
            pass
        self.assertIsNone(json.loads(state.read_text(encoding="utf8"))["lock"])

    def test_export_unknown_rows_creates_nothing(self):
        bank = self.make_bank([("ceiling", FX)])
        out = self.d / "new.xlsx"
        code, msg = run("bank/export_selection.py", bank, "--rows", "500,501", "--name", "x", "--color", "3000K",
                        "--skeleton", self.legend, "--out", out)
        self.assertNotEqual(code, 0)
        self.assertFalse(out.exists(), "an empty legend was left behind")

    def test_export_bad_rows_argument(self):
        bank = self.make_bank([("ceiling", FX)])
        code, msg = run("bank/export_selection.py", bank, "--rows", "9,abc", "--skeleton", self.legend, "--out", self.d / "n.xlsx")
        self.assertNotEqual(code, 0)

    def test_export_partly_missing_rows_warns(self):
        bank = self.make_bank([("ceiling", FX)])
        out = self.d / "new.xlsx"
        code, msg = run("bank/export_selection.py", bank, "--rows", "9,500", "--name", "x", "--color", "3000K",
                        "--skeleton", self.legend, "--out", out)
        self.assertEqual(code, 0)
        self.assertIn("500", msg)
        self.assertTrue(out.exists())


class TestReturnedFile(BankBase):
    def setUp(self):
        super().setUp()
        self.bank = self.make_bank([("ceiling", FX)])

    def add_new(self, data, image=True):
        row = IF.insert_fixture(str(self.legend), "wall", data, image_path=str(self.photo) if image else None)
        code, _ = run("bank/mark_bank_status.py", self.legend, "--row", row, "--status", "new")
        self.assertEqual(code, 0)
        return row

    def bank_count(self):
        ws = openpyxl.load_workbook(self.bank)[M.BANK_SHEET_NAME]
        hr = M.find_bank_header_row(ws)
        return sum(1 for r in range(hr + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value)

    def test_unverified_orange_row_does_not_enter_bank(self):
        self.add_new({**FX, "תיאור": "מנורת קיר מדגם Nova\nLED 6W", "גוון": ""})
        code, out = run("bank/process_new_fixtures.py", self.legend, self.bank, "--no-rebuild")
        self.assertEqual(code, 0)
        self.assertIn("כתום", out)
        self.assertEqual(self.bank_count(), 1)

    def test_row_without_image_is_skipped_with_reason(self):
        self.add_new({**FX, "תיאור": "מנורת קיר מדגם Nova\nLED 6W"}, image=False)
        code, out = run("bank/process_new_fixtures.py", self.legend, self.bank, "--no-rebuild")
        self.assertEqual(self.bank_count(), 1)
        self.assertIn("דולגו", out)

    def test_clean_row_is_added_once(self):
        self.add_new({**FX, "תיאור": "מנורת קיר מדגם Nova\nLED 6W"})
        run("bank/process_new_fixtures.py", self.legend, self.bank, "--no-rebuild")
        run("bank/process_new_fixtures.py", self.legend, self.bank, "--no-rebuild")
        self.assertEqual(self.bank_count(), 2)


# --------------------------------------------------- matching: model numbers, thresholds
B_ = "\nדרייברים למיקום באזור פריק ונגיש"


def desc(kind, name, spec):
    return f"{kind} מדגם {name}\n{spec}{B_}"


class TestMatchRules(BankBase):
    MFR = "HALVARD LUMEN"

    def verdict(self, bank_name, cand_name, bank_spec="LED 8W 700LM 3000K CRI90", cand_spec=None):
        bank = self.make_bank([("floor", {**FX, "יצרן": self.MFR, "תיאור": desc("שקוע רצפה", bank_name, bank_spec)})])
        return MF.find_match(self.MFR, desc("שקוע רצפה", cand_name, cand_spec or bank_spec), bank)

    def test_number_helpers(self):
        d = lambda n: desc("ספוט", n, "LED 8W")
        self.assertTrue(MF.model_numbers_differ(d("Kestra 2642"), d("Kestra 2632")))
        self.assertFalse(MF.model_numbers_differ(d("Arca Mini"), d("Arca Mini")))
        self.assertFalse(MF.model_numbers_differ(d("Arca Mini"), d("Arca Mini 40")))   # only one side has a number
        # "differ" is the weaker rule: an extra part number also counts as different ...
        self.assertTrue(MF.model_numbers_differ(d("Orbo Floor 2"), d("Orbo Floor 2 F7741020")))
        # ... "conflict" is stricter: it needs two lists where neither contains the other
        self.assertFalse(MF.model_numbers_conflict(d("Orbo Floor 2"), d("Orbo Floor 2 F7741020")))
        self.assertTrue(MF.model_numbers_conflict(d("Nexa 35"), d("Nexa 50")))
        self.assertTrue(MF.model_numbers_conflict(d("Kestra 2642"), d("Kestra 2632")))

    def test_different_model_number_is_new_not_maybe(self):
        for a, b in (("Kestra 2642", "Kestra 2632"), ("Nexa 35", "Nexa 50")):
            r = self.verdict(a, b)
            self.assertEqual(r["verdict"], "new", (a, b, r))
            self.assertLess(r["score"], MF.MAYBE_THRESHOLD)

    def test_same_model_is_strong(self):
        self.assertEqual(self.verdict("Arca Mini", "Arca Mini")["verdict"], "strong")

    def test_extra_part_number_is_the_same_product_and_stays_maybe(self):
        """'Orbo Floor 2' and 'Orbo Floor 2 F7741020' are one product: not "new", but never "strong"
        either, so a person confirms the merge."""
        r = self.verdict("Orbo Floor 2", "Orbo Floor 2 F7741020")
        self.assertEqual(r["verdict"], "maybe", r)

    def test_threshold_sits_in_the_gap_between_new_and_similar(self):
        """Real bank, 22.10: new fixtures from a known manufacturer scored 0.51-0.55, genuinely similar
        ones 0.71-0.79. With the old 0.45 cut-off every new fixture was "maybe" and none was ever added.
        Same shape here with invented products."""
        terra = "LED 5W 410LM 3000K CRI90 IP67"
        new = self.verdict("Terra", "Nova", terra, "LED 6W 400LM")
        self.assertEqual(new["verdict"], "new", new)
        self.assertGreater(new["score"], 0.45, "this case must be one the old threshold got wrong")
        self.assertLess(new["score"], MF.MAYBE_THRESHOLD)
        for name, spec in (("Terra Mini", "LED 7W 500LM 4000K CRI80 IP65"),
                           ("Terra Lite", "LED 9W 700LM 4000K CRI80 IP54\nזוית פיזור 25 מעלות")):
            similar = self.verdict("Terra", name, terra, spec)
            self.assertEqual(similar["verdict"], "maybe", (name, similar))
            self.assertGreaterEqual(similar["score"], MF.MAYBE_THRESHOLD)
            self.assertLess(similar["score"], MF.STRONG_THRESHOLD)


# ------------------------------------------- returned files: duplicates and failed saves
NEW_WALL = {**FX, "תיאור": "מנורת קיר מדגם Nova\nLED 6W\nדרייברים למיקום באזור פריק ונגיש"}


class TestReturnedFileEdgeCases(BankBase):
    def setUp(self):
        super().setUp()
        self.bank = self.make_bank([("ceiling", FX)])
        self.legend2 = make_template.build(self.d / "legend2.xlsx")

    def returned(self, legend):
        row = IF.insert_fixture(str(legend), "wall", NEW_WALL, image_path=str(self.photo))
        code, _ = run("bank/mark_bank_status.py", legend, "--row", row, "--status", "new")
        self.assertEqual(code, 0)
        return row

    def bank_count(self):
        ws = openpyxl.load_workbook(self.bank)[M.BANK_SHEET_NAME]
        hr = M.find_bank_header_row(ws)
        return sum(1 for r in range(hr + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value)

    def comment_of(self, legend, row):
        ws = openpyxl.load_workbook(legend)[IF.SHEET_NAME]
        anchor = next(c for c in range(1, ws.max_column + 1) if ws.cell(row=row, column=c).comment)
        return ws.cell(row=row, column=anchor).comment.text

    def test_two_projects_returning_the_same_new_fixture_add_it_once(self):
        """Both reviewers were told "new" against the same bank; the second to return must not duplicate."""
        row_a = self.returned(self.legend)
        row_b = self.returned(self.legend2)
        code, out = run("bank/process_new_fixtures.py", self.legend, self.bank, "--no-rebuild")
        self.assertEqual(code, 0)
        self.assertEqual(self.bank_count(), 2)
        code, out = run("bank/process_new_fixtures.py", self.legend2, self.bank, "--no-rebuild")
        self.assertEqual(code, 0)
        self.assertIn("כבר בבנק", out)
        self.assertEqual(self.bank_count(), 2, "the second project duplicated the fixture")
        # and its legend is updated, so it is not picked up again
        self.assertIn("נוסף", self.comment_of(self.legend2, row_b))

    def test_failed_legend_save_says_the_bank_is_already_updated_and_a_rerun_is_safe(self):
        """The bank is written first, the legend second. If the legend is open in Excel the second step
        fails -- the message must not claim "nothing happened", and running again must not duplicate."""
        import process_new_fixtures as PNF
        row = self.returned(self.legend)
        real = M.save_workbook

        def fail_for_legend(wb, path):
            if Path(path) == Path(self.legend):
                raise M.WorkbookError("הקובץ פתוח באקסל")
            return real(wb, path)

        with mock.patch.object(M, "save_workbook", side_effect=fail_for_legend):
            with self.assertRaises(ValueError) as e:
                PNF.process(str(self.legend), str(self.bank))
        self.assertIn("כבר נכנסו לבנק", str(e.exception))
        self.assertEqual(self.bank_count(), 2)

        code, out = run("bank/process_new_fixtures.py", self.legend, self.bank, "--no-rebuild")   # legend closed now
        self.assertEqual(code, 0)
        self.assertEqual(self.bank_count(), 2, "re-run after the failed save duplicated the fixture")
        self.assertIn("נוסף", self.comment_of(self.legend, row))


# --------------------------------------------------------- bank_guard: backups and locks
class TestBankGuard(BankBase):
    def setUp(self):
        super().setUp()
        self.bank = self.make_bank([("ceiling", FX)])
        self.state = self.bank.parent / ".bank_state.json"

    def lock(self):
        return json.loads(self.state.read_text(encoding="utf8")).get("lock")

    def backups(self):
        return {p.name: sha(p) for p in (self.bank.parent / "backups").glob("*.xlsx")}

    def test_damaged_bank_does_not_overwrite_good_backups(self):
        """Found in QA: a damaged bank + 5 failed runs overwrote all 5 backups with the damaged file."""
        for _ in range(2):
            with bank_guard(self.bank, "healthy run"):
                pass
        good = self.backups()
        self.assertEqual(len(good), 2)

        self.bank.write_bytes(self.bank.read_bytes()[:500])          # a truncated download / half-synced file
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            for _ in range(6):
                with bank_guard(self.bank, "run on a damaged bank"):
                    pass
        self.assertEqual(self.backups(), good, "a backup was overwritten or added from the damaged bank")
        self.assertIn("לא נשמר גיבוי חדש", out.getvalue())

    def test_healthy_bank_is_backed_up_and_rotates(self):
        for _ in range(6):
            with bank_guard(self.bank, "run"):
                pass
        self.assertEqual(sorted(self.backups()), [f"bank-backup-{i}.xlsx" for i in range(1, 6)])

    def test_state_file_is_written_atomically(self):
        import bank_guard as G
        G._write_state(self.state, {"x": 1})
        self.assertEqual(json.loads(self.state.read_text(encoding="utf8")), {"x": 1})
        self.assertEqual(list(self.d.glob(".~.bank_state*")), [], "temp state file left behind")

    def test_crashed_run_lock_expires_without_waiting_ten_minutes(self):
        """kill -9 in the middle of a run: the heartbeat stops, so the lock goes stale after ~2 minutes
        (it used to take 10), and unlock_bank.py can release it by hand."""
        child_code = (
            "import sys, time\n"
            f"sys.path.insert(0, {str(ROOT / 'bank')!r})\n"
            "import bank_guard as G\n"
            "G.HEARTBEAT_SECONDS = 0.2\n"                       # shorten only the test's clock
            f"with G.bank_guard({str(self.bank)!r}, 'child run'):\n"
            "    print('locked', flush=True)\n"
            "    time.sleep(60)\n")
        child = subprocess.Popen([PY, "-c", child_code], stdout=subprocess.PIPE, text=True,
                                 encoding="utf-8", env=ENV)
        try:
            for line in child.stdout:                            # skip the "backup saved" line
                if line.strip() == "locked":
                    break
            else:
                self.fail("the child run never took the lock")
            t1 = self.lock()["ts"]
            time.sleep(0.8)
            t2 = self.lock()["ts"]
            self.assertGreater(t2, t1, "the heartbeat did not refresh the lock while the run was alive")
        finally:
            child.kill()                                         # SIGKILL: no `finally`, no clean-up
            child.wait()
            child.stdout.close()

        t3 = self.lock()["ts"]
        time.sleep(0.8)
        self.assertEqual(self.lock()["ts"], t3, "the lock kept being refreshed after the process died")

        # right after the crash the lock still looks fresh: a second run is stopped, and told how to continue
        with self.assertRaises(SystemExit) as e:
            with bank_guard(self.bank, "second run"):
                pass
        self.assertIn("תפוס", str(e.exception))
        self.assertIn("unlock_bank.py", str(e.exception))

        # unlock_bank.py refuses a fresh lock unless --force
        code, out = run("bank/unlock_bank.py", self.bank)
        self.assertNotEqual(code, 0)
        self.assertIsNotNone(self.lock())
        code, out = run("bank/unlock_bank.py", self.bank, "--force")
        self.assertEqual(code, 0)
        self.assertIsNone(self.lock())

    def test_lock_without_heartbeat_for_two_minutes_is_stale(self):
        st = json.loads(self.state.read_text(encoding="utf8")) if self.state.exists() else {}
        st["lock"] = {"who": "x", "action": "crashed", "ts": time.time() - 121, "id": "dead"}
        self.state.write_text(json.dumps(st), encoding="utf8")
        with bank_guard(self.bank, "after a crash"):
            pass                                                 # not blocked
        self.assertIsNone(self.lock())

    def test_unlock_releases_a_stale_lock_without_force(self):
        self.state.write_text(json.dumps({"lock": {"who": "x", "action": "crashed", "ts": time.time() - 500}}),
                              encoding="utf8")
        code, out = run("bank/unlock_bank.py", self.bank)
        self.assertEqual(code, 0)
        self.assertIsNone(self.lock())

    def test_unlock_on_an_unlocked_bank_is_harmless(self):
        code, out = run("bank/unlock_bank.py", self.bank)
        self.assertEqual(code, 0)
        self.assertIn("לא נעול", out)


# ------------------------------- export: row fingerprints and file names
FX2 = {**FX, "יצרן": "HALVARD LUMEN", "תיאור": "שקוע רצפה מדגם Terra\nLED 5W 410LM 3000K\nדרייברים למיקום באזור פריק ונגיש"}


class TestExportSafety(BankBase):
    def setUp(self):
        super().setUp()
        self.bank = self.make_bank([("ceiling", FX), ("floor", FX2)])
        self.out = self.d / "new.xlsx"

    def bank_rows(self):
        ws = openpyxl.load_workbook(self.bank)[M.BANK_SHEET_NAME]
        hr = M.find_bank_header_row(ws)
        return [(r, M.normalize(ws.cell(row=r, column=M.col_index("יצרן")).value),
                 M.normalize_text(ws.cell(row=r, column=M.col_index("תיאור")).value))
                for r in range(hr + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value]

    def export(self, rows, name="x", out=None):
        return run("bank/export_selection.py", self.bank, "--rows", rows, "--name", name, "--color", "3000K",
                   "--skeleton", self.legend, "--out", out or self.out)

    def sort_the_bank(self):
        """What a person does in Excel with the sort button: rows swap places."""
        wb = openpyxl.load_workbook(self.bank)
        ws = wb[M.BANK_SHEET_NAME]
        (ra, *_), (rb, *_) = self.bank_rows()
        for field in ("יצרן", "תיאור"):
            c = M.col_index(field)
            a, b = ws.cell(row=ra, column=c).value, ws.cell(row=rb, column=c).value
            ws.cell(row=ra, column=c).value, ws.cell(row=rb, column=c).value = b, a
        wb.save(self.bank)

    # --- fingerprints
    def test_fingerprint_is_stable_and_tells_fixtures_apart(self):
        (_, m1, d1), (_, m2, d2) = self.bank_rows()
        fp = M.row_fingerprint(m1, d1)
        self.assertRegex(fp, r"^[0-9a-f]{4}$")
        self.assertEqual(fp, M.row_fingerprint(m1.lower(), "  " + d1.replace("\n", "  ") + " "))   # case / spacing don't matter
        self.assertNotEqual(fp, M.row_fingerprint(m2, d2))

    def test_page_carries_a_fingerprint_for_every_row(self):
        code, out = run("pages/build_pages.py", "--bank", self.bank, "--outdir", self.d)
        self.assertEqual(code, 0)
        html = (self.d / "picker.html").read_text(encoding="utf8")
        for _, man, desc in self.bank_rows():
            self.assertIn(f'"fp": "{M.row_fingerprint(man, desc)}"', html)

    def test_export_with_matching_fingerprints_works(self):
        rows = ",".join(f"{r}:{M.row_fingerprint(m, d)}" for r, m, d in self.bank_rows())
        code, out = self.export(rows)
        self.assertEqual(code, 0, out)
        self.assertTrue(self.out.exists())

    def test_plain_row_numbers_still_work(self):
        code, out = self.export(",".join(str(r) for r, _, _ in self.bank_rows()))
        self.assertEqual(code, 0, out)

    def test_sorted_bank_stops_the_export_instead_of_exporting_another_fixture(self):
        """Found in QA: after the bank was sorted in Excel, an old page's "row 31" was a different fixture
        (another manufacturer's fixture) and the wrong one was exported without a word."""
        (ra, ma, da), _ = self.bank_rows()
        asked = f"{ra}:{M.row_fingerprint(ma, da)}"          # what the page saw before the sort
        self.sort_the_bank()
        before = sha(self.bank)
        code, out = self.export(asked)
        self.assertNotEqual(code, 0)
        self.assertIn("הבנק השתנה", out)
        self.assertFalse(self.out.exists(), "a legend with the wrong fixture was created")
        self.assertEqual(sha(self.bank), before)
        self.assertNotIn("Traceback", out)

    # --- file names
    def test_forbidden_characters_in_a_project_name_do_not_break_the_file_name(self):
        """Found in QA: a "/" in the name created a folder instead of part of the file name."""
        rows = ",".join(str(r) for r, _, _ in self.bank_rows())
        code, out = run("bank/export_selection.py", self.bank, "--rows", rows, "--name", 'בית/צפוני: "א"',
                        "--color", "3000K", "--skeleton", self.legend)
        self.assertEqual(code, 0, out)
        self.assertIn("הוחלפו", out)
        made = [p for p in self.d.glob("מקרא - *.xlsx")]
        self.assertEqual(len(made), 1, made)
        self.assertFalse(set('\\/:*?"<>|') & set(made[0].name))
        self.assertEqual([p for p in self.d.iterdir() if p.is_dir() and p.name != "backups"], [], "a stray folder was created")

    def test_name_that_is_empty_once_cleaned_creates_nothing(self):
        """"..." is not a usable Windows file name (trailing dots are dropped), so nothing is created."""
        rows = ",".join(str(r) for r, _, _ in self.bank_rows())
        code, out = run("bank/export_selection.py", self.bank, "--rows", rows, "--name", "...",
                        "--color", "3000K", "--skeleton", self.legend)
        self.assertNotEqual(code, 0)
        self.assertEqual(list(self.d.glob("מקרא - *.xlsx")), [])


# ------------------------------------- legend: skeleton row and leftover temp files
class TestSkeletonAndTempFiles(Tmp):
    def descriptions(self):
        ws = self.sheet()
        col = IF.build_header_map(ws)["תיאור"]
        return [str(ws.cell(row=r, column=col).value or "") for r in range(1, ws.max_row + 1)]

    def test_skeleton_text_helper(self):
        self.assertTrue(IF._is_skeleton_text("[סוג הגוף] מדגם [שם הדגם]\nLED [הספק]W"))
        self.assertFalse(IF._is_skeleton_text("[Ø55] שקוע תקרה מדגם Arca Mini"))
        self.assertFalse(IF._is_skeleton_text("[IP65] פנס חוץ מדגם Terra"))

    def test_real_description_starting_with_a_bracket_is_kept(self):
        """Found in QA: the example row was recognised by "starts with [", so a real fixture whose
        description began with "[Ø55]" was deleted together with it."""
        data = {**FX, "תיאור": "[Ø55] שקוע תקרה מדגם Arca\nLED 7W 640LM 3000K"}
        IF.insert_fixture(str(self.legend), "ceiling", data, image_path=str(self.photo))
        IF.insert_fixture(str(self.legend), "wall", FX, image_path=str(self.photo))   # strips again on every save
        texts = self.descriptions()
        self.assertTrue(any(t.startswith("[Ø55]") for t in texts), "the bracketed fixture was deleted")
        self.assertFalse(any("[סוג הגוף]" in t for t in texts), "the template example row is still there")

    def test_repeated_crashes_leave_at_most_one_temp_file(self):
        """Hard-kill the process at the moment of the swap, three times. The temp name is fixed per target,
        so crashes don't pile up orphans, and the original is never touched."""
        before = sha(self.legend)
        child = ("import os, sys\n"
                 f"sys.path.insert(0, {str(ROOT / 'skill' / 'scripts')!r})\n"
                 "import xlsx_io, insert_fixture as IF\n"
                 "os.replace = lambda *a, **k: os._exit(9)\n"          # dies after the temp file is fully written
                 f"IF.insert_fixture({str(self.legend)!r}, 'ceiling', {FX!r})\n")
        for _ in range(3):
            r = subprocess.run([PY, "-c", child], capture_output=True, env=ENV)
            self.assertEqual(r.returncode, 9, r.stderr)
        self.assertEqual(len(list(self.d.glob(".~saving-*"))), 1)
        self.assertEqual(sha(self.legend), before)

        IF.insert_fixture(str(self.legend), "ceiling", FX, image_path=str(self.photo))   # a normal run afterwards
        self.assertEqual(list(self.d.glob(".~saving-*")), [], "the leftover was not cleaned up by the next save")
        self.assertEqual(self.cell(self.sheet(), 10, "יצרן").value, "VELTRANO")


# ------------------------------------------- one-off migration: restore line breaks
FX3 = {**FX, "יצרן": "SOLVIK LUCE", "תיאור": "מנורת קיר מדגם Nova\nLED 6W 400LM 3000K\nדרייברים למיקום באזור פריק ונגיש",
       "מתח_עמעום": "24V\nDIM DALI"}
MULTILINE = ("תיאור", "מתח_עמעום")


class TestRestoreLineBreaks(BankBase):
    """The import used to flatten the bank's descriptions to one line; this migration puts the line breaks
    back from the old legends, without ever changing a word."""
    SCRIPT = "bank/migrations/restore_line_breaks.py"

    def setUp(self):
        super().setUp()
        self.bank = self.make_bank([("ceiling", FX), ("floor", FX2), ("wall", FX3)])
        self.src = self.d / "old-legends"
        self.src.mkdir()
        shutil.copy(self.d / "old.xlsx", self.src / "old.xlsx")      # the legend the bank was built from
        self.expected = self.cells()                                 # texts WITH line breaks
        self.flatten()

    def cells(self):
        ws = openpyxl.load_workbook(self.bank)[M.BANK_SHEET_NAME]
        hr = M.find_bank_header_row(ws)
        return {r: {c: ws.cell(row=r, column=M.col_index(c)).value for c, _ in M.BANK_COLUMNS}
                for r in range(hr + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value}

    def flatten(self, edit=None):
        """Recreate what the old import did to the bank: every line break becomes a space."""
        wb = openpyxl.load_workbook(self.bank)
        ws = wb[M.BANK_SHEET_NAME]
        for r in self.expected:
            for c in MULTILINE:
                cell = ws.cell(row=r, column=M.col_index(c))
                cell.value = M.normalize(cell.value)
        if edit:
            edit(ws)
        wb.save(self.bank)
        self.flat = self.cells()

    def test_sanity_the_test_bank_is_really_flat(self):
        for r, row in self.flat.items():
            self.assertNotIn("\n", row["תיאור"])

    def test_line_breaks_come_back_and_no_word_changes(self):
        code, out = run(self.SCRIPT, self.bank, "--sources", self.src)
        self.assertEqual(code, 0, out)
        after = self.cells()
        for r, row in after.items():
            for c in MULTILINE:
                self.assertEqual(row[c], M.normalize_text(self.expected[r][c]), (r, c))
            for c, v in row.items():                                  # every other column is untouched
                if c not in MULTILINE:
                    self.assertEqual(v, self.flat[r][c], (r, c))
        self.assertIn("ישוחזרו: 6", out)                              # 3 descriptions + 3 voltage cells

    def test_dry_run_changes_nothing(self):
        before = sha(self.bank)
        code, out = run(self.SCRIPT, self.bank, "--sources", self.src, "--dry-run")
        self.assertEqual(code, 0, out)
        self.assertEqual(sha(self.bank), before)
        self.assertIn("ישוחזרו: 6", out)

    def test_in_place_run_goes_through_lock_and_backup(self):
        run(self.SCRIPT, self.bank, "--sources", self.src)
        self.assertTrue((self.bank.parent / "backups" / "bank-backup-1.xlsx").exists())
        self.assertIsNone(json.loads((self.bank.parent / ".bank_state.json").read_text(encoding="utf8"))["lock"])

    def test_out_writes_a_copy_and_leaves_the_bank_alone(self):
        before = sha(self.bank)
        out_file = self.d / "restored.xlsx"
        code, _ = run(self.SCRIPT, self.bank, "--sources", self.src, "--out", out_file)
        self.assertEqual(code, 0)
        self.assertEqual(sha(self.bank), before)
        self.assertTrue(out_file.exists())

    def test_two_different_versions_in_the_sources_are_left_alone(self):
        """If the old legends disagree about where the lines break, there is no right answer: don't guess."""
        (r_floor, _) = [(r, row) for r, row in self.flat.items() if "Terra" in row["תיאור"]][0]
        other = openpyxl.Workbook()
        other.active["A1"] = FX2["תיאור"].replace("\nLED", " LED", 1)       # same words, different breaks
        other.save(self.src / "other-version.xlsx")
        code, out = run(self.SCRIPT, self.bank, "--sources", self.src)
        self.assertEqual(code, 0, out)
        after = self.cells()
        self.assertNotIn("\n", after[r_floor]["תיאור"])                     # untouched
        self.assertIn("(לא נגעתי): 1", out)
        others = [r for r in after if r != r_floor]
        self.assertTrue(all("\n" in after[r]["תיאור"] for r in others))      # the unambiguous rows were restored

    def test_text_that_is_not_in_the_sources_is_left_alone(self):
        """A description a person edited in the bank since the import has no source to restore from."""
        def edit(ws):
            r = next(iter(self.expected))
            c = M.col_index("תיאור")
            ws.cell(row=r, column=c).value = ws.cell(row=r, column=c).value + " (נערך ידנית)"
        self.flatten(edit)
        code, out = run(self.SCRIPT, self.bank, "--sources", self.src)
        self.assertEqual(code, 0, out)
        edited_row = next(iter(self.expected))
        self.assertNotIn("\n", self.cells()[edited_row]["תיאור"])
        self.assertIn("לא נמצאו במקור: 1", out)

    def test_sources_are_required(self):
        code, out = run(self.SCRIPT, self.bank)
        self.assertNotEqual(code, 0)
        self.assertNotIn("Traceback", out)


# ------------------------------------------------------------------ the web page
class TestPage(BankBase):
    def test_description_cannot_break_out_of_the_script_tag(self):
        bank = self.make_bank([("ceiling", {**FX, "תיאור": "ספוט מדגם X</script><b>\nLED 1W"})])
        code, out = run("pages/build_pages.py", "--bank", bank, "--outdir", self.d)
        self.assertEqual(code, 0)
        html = (self.d / "picker.html").read_text(encoding="utf8")
        self.assertNotIn("X</script>", html)

    def test_not_a_bank_file(self):
        code, out = run("pages/build_pages.py", "--bank", self.legend, "--outdir", self.d)
        self.assertNotEqual(code, 0)
        self.assertNotIn("Traceback", out)


if __name__ == "__main__":
    unittest.main()
