#!/usr/bin/env python3
"""
Scan a reviewed/corrected project מקרא workbook for rows still carrying the
"🆕 גוף חדש" marker (put there by mark_bank_status.py right after mikra
inserted them — see mikra/SKILL.md step 3.5), add each one to the fixture
bank, and flip its marker to "✓ נוסף לבנק" so a second run over the same
returned file never adds it twice.

This is the step that runs when the employee sends the corrected legend file
back (mikra SKILL.md step 4) — the act of returning the reviewed file IS the
approval; no separate chat confirmation is required. A row the employee
decided doesn't belong in the bank should have had its comment deleted before
sending the file back (the marker's own text says so) — this script simply
does nothing for a row with no "🆕" comment.

Usage:
    python3 process_new_fixtures.py <מקרא.xlsx> <בנק.xlsx> \
        [--mikra-out /path] [--bank-out /path]

Without --mikra-out / --bank-out, both files are updated in place.
"""
import argparse
from bank_guard import bank_guard
from rebuild_pages import rebuild
import json
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.cell.rich_text import CellRichText

sys.path.insert(0, str(Path(__file__).parent))
import migrate_mikra_to_bank as M  # noqa: E402
import mark_bank_status as MK  # noqa: E402
from add_fixture_to_bank import add_fixture  # noqa: E402

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")



ORANGE = {"FFFFC000", "00FFC000", "FFC000"}


def still_flagged(ws, row, header_map):
    """Fields in this row that are still marked orange (whole-cell fill, or an
    orange phrase inside the text). The returned file is the approval -- but an
    orange cell means nobody confirmed that value yet, and the bank is shared:
    an unverified value there would be copied into every future project."""
    flagged = []
    for field, col in header_map.items():
        cell = ws.cell(row=row, column=col)
        rgb = str(getattr(cell.fill.fgColor, "rgb", "") or "").upper()
        if cell.fill.fill_type == "solid" and rgb in ORANGE:
            flagged.append(field)
            continue
        v = cell.value
        if isinstance(v, CellRichText):
            for part in v:
                font = getattr(part, "font", None)
                colour = str(getattr(getattr(font, "color", None), "rgb", "") or "").upper()
                if colour in ORANGE:
                    flagged.append(field)
                    break
    return flagged


def process(mikra_path, bank_path, mikra_out=None, bank_out=None):
    mikra_out = mikra_out or mikra_path
    bank_out = bank_out or bank_path
    today = date.today().strftime("%d.%m.%Y")

    # --- Pass 1: read-only. Gather every "🆕" row's fixture data (including
    # image bytes) from a workbook we never save. openpyxl images are backed
    # by a lazily-opened stream from the source zip; draining that stream via
    # img._data() and then saving THIS SAME workbook object again later makes
    # openpyxl try to re-read that now-closed stream and crash ("I/O
    # operation on closed file") — so the object that had its images read is
    # discarded, never re-saved. Pass 2 below re-opens the file fresh instead.
    wb_read = M.open_workbook(mikra_path, rich_text=True)
    if MK.SHEET_NAME not in wb_read.sheetnames:
        raise ValueError(f"No '{MK.SHEET_NAME}' sheet in {mikra_path}.")
    ws_read = wb_read[MK.SHEET_NAME]

    header_row, header_map = M.find_header_row(ws_read)
    if header_row is None:
        raise ValueError("Could not find a header row with both יצרן and תיאור in the first 25 rows.")
    blocks = M.find_category_blocks(ws_read, header_row=header_row)
    img_map = M.images_by_anchor_row(ws_read)
    anchor_col = MK.find_anchor_col(ws_read, header_row)
    project_name = M.project_name_from_header(ws_read)

    candidates, skipped = [], []

    for category, (first, last) in blocks.items():
        row = first
        while row <= last:
            cell = ws_read.cell(row=row, column=anchor_col)
            if isinstance(cell, MergedCell) or cell.value in (None, ""):
                row += 1
                continue

            comment = cell.comment
            if comment and str(comment.text).startswith(MK.NEW_PREFIX):
                def val(field):
                    col = header_map.get(field)
                    if not col:
                        return ""
                    v = ws_read.cell(row=row, column=col).value
                    return M.normalize_text(v) if field in M.MULTILINE_FIELDS else M.normalize(v)

                yatzran, tiur = val("יצרן"), val("תיאור")
                img = img_map.get(row - 1)  # 0-indexed anchor

                flagged = still_flagged(ws_read, row, header_map)
                if not yatzran or not tiur:
                    skipped.append({"row": row, "reason": "חסר יצרן או תיאור"})
                elif flagged:
                    skipped.append({"row": row, "reason": "עדיין מסומן בכתום (לא נבדק): " + ", ".join(flagged)
                                    + " — לבדוק, להסיר את הצביעה ולהחזיר שוב"})
                elif img is None:
                    skipped.append({"row": row, "reason": "לא נמצאה תמונה מתאימה"})
                else:
                    candidates.append({
                        "row": row, "category": category, "יצרן": yatzran, "תיאור": tiur,
                        "גוון": val("גוון"), "מתח_עמעום": val("מתח_עמעום"),
                        "סימון": val("סימון"), "project_name": project_name,
                        "image_bytes": img._data(), "image_w": img.width, "image_h": img.height,
                    })

            span = 1
            for mc in ws_read.merged_cells.ranges:
                if mc.min_col <= anchor_col <= mc.max_col and mc.min_row == row:
                    span = max(span, mc.max_row - mc.min_row + 1)
            row += span

    del wb_read, ws_read  # done reading; never saved, never re-touched

    # --- Pass 2: add each candidate to the bank.
    added = []
    current_bank_path = bank_path
    for fixture in candidates:
        bank_row = add_fixture(current_bank_path, bank_out, fixture)
        current_bank_path = bank_out  # subsequent adds build on the just-saved bank
        added.append({"row": fixture["row"], "bank_row": bank_row,
                       "יצרן": fixture["יצרן"], "תיאור": fixture["תיאור"]})

    # --- Pass 3: flip each added row's marker to "added", from a FRESH load
    # (comments only — no image access, so no re-save conflict).
    if added:
        wb_write = M.open_workbook(mikra_path)
        ws_write = wb_write[MK.SHEET_NAME]
        for f in added:
            cell = ws_write.cell(row=f["row"], column=anchor_col)
            cell.comment = Comment(MK.comment_text("added", date_str=today), MK.COMMENT_AUTHOR)
        Path(mikra_out).parent.mkdir(parents=True, exist_ok=True)
        M.save_workbook(wb_write, mikra_out)
    elif mikra_out != mikra_path:
        import shutil
        Path(mikra_out).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(mikra_path, mikra_out)

    return {
        "project_name": project_name, "added": added, "skipped": skipped,
        "mikra_out": mikra_out, "bank_out": current_bank_path if added else bank_path,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mikra_xlsx")
    ap.add_argument("bank_xlsx")
    ap.add_argument("--mikra-out", default=None)
    ap.add_argument("--bank-out", default=None)
    ap.add_argument("--no-rebuild", action="store_true",
                    help="לא לבנות מחדש את דפי העבודה בסוף (ברירת מחדל: כן, אם נוספו גופים)")
    args = ap.parse_args()

    try:
        with bank_guard(args.bank_xlsx, "הוספת גופים חדשים מקובץ מקרא שחזר"):
            report = process(args.mikra_xlsx, args.bank_xlsx, args.mikra_out, args.bank_out)
    except (M.WorkbookError, ValueError) as e:
        raise SystemExit(f"שגיאה: {e}")
    print(f"פרויקט: {report['project_name']}")
    print(f"נוספו {len(report['added'])} גופים חדשים לבנק:")
    for f in report["added"]:
        print(f"  + row {f['row']} -> bank row {f['bank_row']}: {f['יצרן']} — {f['תיאור'][:50].replace(chr(10), ' · ')}")
    if report["skipped"]:
        print(f"דולגו {len(report['skipped'])} שורות מסומנות (לא ניתן להוסיף):")
        for s in report["skipped"]:
            print(f"  - row {s['row']}: {s['reason']}")
    print(f"\nמקרא נשמר: {report['mikra_out']}")
    print(f"בנק נשמר: {report['bank_out']}")

    if report["added"] and not args.no_rebuild:
        rebuild(report["bank_out"])


if __name__ == "__main__":
    main()
