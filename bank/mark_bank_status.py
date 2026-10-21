#!/usr/bin/env python3
"""
Attach a bank-status comment to one fixture row already written by
mikra/scripts/insert_fixture.py, based on the verdict from match_fixture.py.
Kept as a separate small script (rather than folded into insert_fixture.py)
so the bank-matching step can change independently of the tested
row-writing/highlighting logic there.

The comment goes on the row's anchor cell (סעיף if the sheet has that column,
otherwise יצרן) — never on תיאור/גוון/etc, so it can never collide with an
"uncertain field" comment insert_fixture.py may have already put there.

Usage:
    python3 mark_bank_status.py <xlsx_path> --row 14 --status new
    python3 mark_bank_status.py <xlsx_path> --row 14 --status maybe \
        --bank-description "..."
    python3 mark_bank_status.py <xlsx_path> --row 14 --status added \
        --date 23.08.2026

'new'   — not found in the bank; will be added automatically once the
          corrected legend file is returned (see mikra SKILL.md step 4).
'maybe' — a similar-looking row already exists in the bank; needs a human
          to check and merge/ignore, never auto-actioned.
'added' — written after process_new_fixtures.py has actually added the row
          to the bank, so re-running it on the same file is a no-op instead
          of adding the fixture twice.
"""
import argparse
import re
import sys
from pathlib import Path

from openpyxl.comments import Comment

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "skill" / "scripts"))
from xlsx_io import WorkbookError, open_workbook, save_workbook  # noqa: E402

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")


SHEET_NAME = "אומדן"
COMMENT_AUTHOR = "mikra skill"

NEW_PREFIX = "🆕 גוף חדש"
MAYBE_PREFIX = "⚠ בדיקת בנק"
ADDED_PREFIX = "✓ נוסף לבנק"


def normalize(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def find_header_row(ws, search_rows=25):
    for row in range(1, search_rows + 1):
        found = {normalize(ws.cell(row=row, column=c).value) for c in range(1, ws.max_column + 1)}
        if {"יצרן", "תיאור"} <= found:
            return row
    return None


def find_anchor_col(ws, header_row):
    seif_col = None
    yatzran_col = None
    for col in range(1, ws.max_column + 1):
        raw = normalize(ws.cell(row=header_row, column=col).value)
        if raw.startswith("סעיף") and seif_col is None:
            seif_col = col
        if raw == "יצרן" and yatzran_col is None:
            yatzran_col = col
    return seif_col or yatzran_col or 1


def comment_text(status, bank_description=None, date_str=None):
    if status == "new":
        return (f"{NEW_PREFIX} — לא נמצא בבנק גופי התאורה. יתווסף לבנק אוטומטית כשהקובץ "
                "המתוקן יוחזר. אם הגוף הזה לא באמת שייך לבנק (חד-פעמי לפרויקט הזה), "
                "מחקו הערה זו לפני ההחזרה.")
    if status == "maybe":
        note = f"{MAYBE_PREFIX}: יתכן שגוף זה כבר קיים בבנק בניסוח דומה"
        if bank_description:
            snippet = bank_description[:80] + ("…" if len(bank_description) > 80 else "")
            note += f' ("{snippet}")'
        return note + " — כדאי לבדוק ולמזג אם כן, כדי למנוע כפילות."
    if status == "added":
        d = date_str or ""
        return f"{ADDED_PREFIX}{(' ב-' + d) if d else ''}."
    raise ValueError("status must be 'new', 'maybe', or 'added'")


def mark(xlsx_path, row, status, bank_description=None, date_str=None, out_path=None):
    wb = open_workbook(xlsx_path, sheet=SHEET_NAME)
    ws = wb[SHEET_NAME]
    header_row = find_header_row(ws)
    if header_row is None:
        raise ValueError("Could not find a header row with both יצרן and תיאור in the first 25 rows.")
    col = find_anchor_col(ws, header_row)
    cell = ws.cell(row=row, column=col)
    cell.comment = Comment(comment_text(status, bank_description, date_str), COMMENT_AUTHOR)
    out_path = out_path or xlsx_path
    save_workbook(wb, out_path)
    print(f"Marked row {row} as '{status}'. Saved to {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx_path")
    ap.add_argument("--row", type=int, required=True)
    ap.add_argument("--status", required=True, choices=["new", "maybe", "added"])
    ap.add_argument("--bank-description", default=None)
    ap.add_argument("--date", default=None)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    mark(args.xlsx_path, args.row, args.status, args.bank_description, args.date, args.out)


if __name__ == "__main__":
    try:
        main()
    except (WorkbookError, ValueError) as e:
        raise SystemExit(f"שגיאה: {e}")
