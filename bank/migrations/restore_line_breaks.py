#!/usr/bin/env python3
"""
ONE-OFF migration: restore line breaks in bank descriptions that were stored flattened.

Background: until the fix for "line breaks flattened in the bank", the import step ran a
whitespace clean-up that was meant for *comparing* texts on the text it *stored* as well.
Every description and voltage cell in the bank therefore became one long line. The code is
fixed (comparison and storage use different clean-ups), but rows imported before the fix
stay flat. A fixture exported from such a row no longer matches the layout of a row written
straight from a PDF (one spec fragment per line).

This script repairs the data in place instead of rebuilding the bank (a rebuild would also
throw away merges and corrections a person made by hand since the import):

  * for every flat bank cell it looks for the SAME text in the source legends (the old
    project files the bank was built from), compared after collapsing whitespace;
  * if exactly one line-broken version exists there, that version is written to the cell;
  * if several different versions exist, or none, the cell is left alone.

What it never does: change a single word (checked per cell before writing), touch any
other column, or touch a row it is not sure about.

Usage:
    python bank/migrations/restore_line_breaks.py <bank.xlsx> --sources <dir> [<dir> ...] \
        [--out <file>] [--dry-run]

  --sources   folder(s) holding the old legends; searched recursively for *.xlsx
  --dry-run   only print what would change
  --out       write the result to another file instead of the bank (no lock / backup then)

Without --out or --dry-run the bank is updated in place, through bank_guard (lock + backup).
"""
import argparse
import sys
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import migrate_mikra_to_bank as M  # noqa: E402
from bank_guard import bank_guard  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")


def collect_sources(dirs):
    """flat text -> set of line-broken versions, from every cell of every source file."""
    found, files, bad = {}, 0, []
    for d in dirs:
        for f in sorted(Path(d).rglob("*.xlsx")):
            if f.name.startswith(("~$", ".~")):
                continue
            try:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            except Exception as e:
                bad.append(f"{f.name} ({e.__class__.__name__})")
                continue
            files += 1
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str) and "\n" in v:
                            found.setdefault(M.normalize(v), set()).add(M.normalize_text(v))
            wb.close()
    return found, files, bad


def plan(bank_path, sources):
    wb = M.open_workbook(bank_path, sheet=M.BANK_SHEET_NAME)
    ws = wb[M.BANK_SHEET_NAME]
    header_row = M.find_bank_header_row(ws)
    changes, stats = [], {"restored": 0, "already": 0, "not_found": 0, "ambiguous": 0}
    for r in range(header_row + 1, ws.max_row + 1):
        for field in M.MULTILINE_FIELDS:
            v = ws.cell(row=r, column=M.col_index(field)).value
            if not isinstance(v, str) or not v.strip():
                continue
            if "\n" in v:
                stats["already"] += 1
                continue
            cands = sources.get(M.normalize(v))
            if cands == {v.strip()}:      # the source had only stray empty lines -> nothing to restore
                stats["not_found"] += 1
                continue
            if not cands:
                stats["not_found"] += 1
            elif len(cands) > 1:
                stats["ambiguous"] += 1
            else:
                new = next(iter(cands))
                assert M.normalize(new) == M.normalize(v)   # same words, only line breaks differ
                changes.append((r, field, v, new))
                stats["restored"] += 1
    return wb, ws, changes, stats


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("bank_xlsx")
    ap.add_argument("--sources", nargs="+", required=True, help="folder(s) with the old legends")
    ap.add_argument("--out", default=None)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args(argv)

    sources, n_files, bad = collect_sources(args.sources)
    print(f"נסרקו {n_files} קבצי מקור" + (f" (לא נפתחו: {', '.join(bad)})" if bad else ""))

    def run():
        wb, ws, changes, stats = plan(args.bank_xlsx, sources)
        print(f"ישוחזרו: {stats['restored']} | כבר עם שבירות: {stats['already']} | "
              f"לא נמצאו במקור: {stats['not_found']} | כמה גרסאות במקור (לא נגעתי): {stats['ambiguous']}")
        if args.dry_run or not changes:
            return changes
        for r, field, _, new in changes:
            cell = ws.cell(row=r, column=M.col_index(field))
            cell.value = new
            # Excel shows a line break only when the cell wraps text
            if not cell.alignment.wrap_text:
                a = cell.alignment
                cell.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical,
                                           wrap_text=True, readingOrder=a.readingOrder)
        M.save_workbook(wb, args.out or args.bank_xlsx)
        print(f"נשמר: {args.out or args.bank_xlsx}")
        return changes

    if args.dry_run or args.out:
        changes = run()
    else:
        with bank_guard(args.bank_xlsx, "שחזור שבירות שורה בתיאורים"):
            changes = run()
    for r, field, old, new in changes[:3]:
        print(f"\nשורה {r} ({field}):\n  לפני: {old}\n  אחרי: " + new.replace("\n", "\n         "))


if __name__ == "__main__":
    try:
        main()
    except (M.WorkbookError, ValueError) as e:
        raise SystemExit(f"שגיאה: {e}")
