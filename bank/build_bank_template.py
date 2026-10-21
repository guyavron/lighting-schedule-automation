#!/usr/bin/env python3
"""
Create an empty בנק גופי תאורה (fixture bank) workbook — a flat, filterable
table (not a visual מקרא-style sheet: no category blocks/colour bars), but
with the same branded header block as מקרא (studio logo + title) above the
column-header row, per Guy's request (23.08.2026) that the bank share
מקרא's visual identity. One row per fixture, AutoFilter on every column so
the studio can filter by ספק / יצרן / קטגוריה / סטטוס etc. directly in Excel.

Columns (fixed order — canonical field names used by migrate_mikra_to_bank.py
to know which column is which, same "read the header row, don't hardcode a
letter" principle as the mikra skill):
    קטגוריה, יצרן, ספק, סטטוס, תיאור, גוון, מתח / פרוטוקול עמעום, סימון,
    תמונה, פרויקט/ים מקור, תאריך הוספה לבנק, בחירה לפרויקט

Header layout mirrors מקרא's own: rows 1-7 reserved for the branded header
block, column-header row at row 8, data from row 9 -- same row number
(find_bank_header_row in migrate_mikra_to_bank.py locates it dynamically
either way, this is just for a fresh build to look right immediately).

Usage:
    python3 build_bank_template.py <out_xlsx_path> [--logo /path/to/logo.png]
"""
import argparse
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")


SHEET_NAME = "בנק"

# Canonical name -> (header text, column width)
COLUMNS = [
    ("קטגוריה", "קטגוריה", 14),
    ("יצרן", "יצרן", 18),
    ("ספק", "ספק", 16),
    ("סטטוס", "סטטוס", 14),
    ("תיאור", "תיאור", 45),
    ("גוון", "גוון", 14),
    ("מתח_עמעום", "מתח / פרוטוקול עמעום", 20),
    ("סימון", "סימון (מהפרויקט המקורי)", 20),
    ("תמונה", "תמונה", 20),
    ("פרויקטים_מקור", "פרויקט/ים מקור", 30),
    ("תאריך_הוספה", "תאריך הוספה לבנק", 16),
    ("בחירה", "בחירה לפרויקט", 16),
]

HEADER_ROW = 8  # matches מקרא's own HEADER_ROW (mikra/scripts/insert_fixture.py)


def build(out_path, logo_path=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for idx, (canon, header_text, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=idx, value=header_text)
        cell.font = Font(name="Arial", size=10)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width

    for r in range(1, HEADER_ROW + 1):
        ws.row_dimensions[r].height = 12.75

    title_cell = ws.cell(row=2, column=2, value="בנק גופי תאורה")
    title_cell.font = Font(name="Arial", size=14, bold=True)
    title_cell.alignment = Alignment(horizontal="right", vertical="center")

    if logo_path:
        img = XLImage(logo_path)
        anchor = f"{get_column_letter(len(COLUMNS) + 1)}1"
        ws.add_image(img, anchor)

    ws.freeze_panes = f"A{HEADER_ROW + 1}"
    last_col_letter = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col_letter}{HEADER_ROW}"

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Created empty bank workbook: {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out_xlsx_path")
    ap.add_argument("--logo", default=None, help="optional logo image to embed in the header")
    args = ap.parse_args()
    build(args.out_xlsx_path, args.logo)


if __name__ == "__main__":
    main()
