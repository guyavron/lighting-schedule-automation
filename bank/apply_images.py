#!/usr/bin/env python3
"""
Replace product photos in the bank, from an explicit row -> image-file map.

Runs only after a human has looked at the photos: some rows imported from old
legend files carry a technical drawing or a screenshot of a catalogue page
instead of a product photo, and a high text-match score doesn't mean the
replacement image is right -- so a person picks, this script applies.

    python bank/apply_images.py --map "12=photos/new.jpg,14=photos/other.jpg" [--dry-run]

Every write goes through bank_guard (lock + backup); pages are rebuilt after.
"""
import argparse, io, json, sys
from contextlib import nullcontext
from bank_guard import bank_guard
from rebuild_pages import rebuild
from pathlib import Path

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "skill" / "scripts"))
from xlsx_io import WorkbookError, open_workbook, save_workbook  # noqa: E402
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
import migrate_mikra_to_bank as M

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")


PROJECT = Path(__file__).resolve().parents[1]


def parse_map(args):
    if args.map_file:
        raw = json.loads(Path(args.map_file).read_text(encoding="utf8"))
        return {int(k): v for k, v in raw.items()}
    out = {}
    for pair in args.map.split(","):
        pair = pair.strip()
        if not pair:
            continue
        if "=" not in pair:
            raise SystemExit(f"פורמט לא תקין: {pair!r} — צפוי שורה=נתיב")
        row, path = pair.split("=", 1)
        out[int(row)] = path.strip()
    return out


def main():
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--map", help='"29=/path/a.jpg,97=/path/a.jpg"')
    g.add_argument("--map-file", help="קובץ JSON: {שורה: נתיב}")
    ap.add_argument("--bank", default=str(PROJECT / "bank.xlsx"))
    ap.add_argument("--dry-run", action="store_true", help="להראות מה יקרה בלי לשמור")
    args = ap.parse_args()

    mapping = parse_map(args)
    missing = [p for p in mapping.values() if not Path(p).exists()]
    if missing:
        raise SystemExit("קבצי תמונה שלא נמצאו:\n  " + "\n  ".join(missing))

    guard = nullcontext() if args.dry_run else bank_guard(args.bank, "החלפת תמונות מוצר")
    with guard:
        run(args, mapping)


def run(args, mapping):
    wb = open_workbook(args.bank, sheet=M.BANK_SHEET_NAME)
    ws = wb[M.BANK_SHEET_NAME]
    header = M.find_bank_header_row(ws)
    for row, path in sorted(mapping.items()):
        if row <= header or row > ws.max_row:
            raise SystemExit(f"שורה {row} מחוץ לטווח שורות הגופים בבנק.")
        man = M.normalize(ws.cell(row=row, column=M.col_index("יצרן")).value)
        desc = M.normalize(ws.cell(row=row, column=M.col_index("תיאור")).value)
        data = Path(path).read_bytes()
        im = Image.open(io.BytesIO(data))
        print(f"שורה {row}: {man} — {desc[:55]}")
        print(f"   <- {Path(path).name} {im.size}")
        if not args.dry_run:
            M.remove_image_at_row(ws, row)
            M.write_image(ws, row, data, im.width, im.height)

    if args.dry_run:
        print("\n--dry-run: לא נשמר כלום.")
        return
    save_workbook(wb, args.bank)
    print(f"\nנשמר: {args.bank}")
    rebuild(args.bank)


if __name__ == "__main__":
    try:
        main()
    except (M.WorkbookError, ValueError) as e:
        raise SystemExit(f"שגיאה: {e}")
