#!/usr/bin/env python3
"""
Append ONE already-confirmed-new fixture to the fixture bank.

This is deliberately the narrow single-row counterpart to
migrate_mikra_to_bank.py's bulk migration: the caller (process_new_fixtures.py)
has already decided, via match_fixture.py + a human's review-and-return of the
legend file, that this exact fixture belongs in the bank — no dedup check runs
here. Reuses migrate_mikra_to_bank.py's own column layout and image-writing
helpers so a bank row added this way is structurally identical to one added by
the bulk migration path.

Usage as a library (see process_new_fixtures.py):
    from add_fixture_to_bank import add_fixture
    add_fixture(bank_path, out_path, {
        "category": "תקרה", "יצרן": "...", "תיאור": "...", "גוון": "...",
        "מתח_עמעום": "...", "סימון": "...", "project_name": "...",
        "image_bytes": b"...", "image_w": 800, "image_h": 600,
    }, supplier="")
"""
import sys
from datetime import date
from pathlib import Path

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "skill" / "scripts"))
from xlsx_io import WorkbookError, open_workbook, save_workbook  # noqa: E402
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).parent))
import migrate_mikra_to_bank as M  # noqa: E402

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")



def add_fixture(bank_path, out_path, fixture, supplier=""):
    wb = open_workbook(bank_path, sheet=M.BANK_SHEET_NAME)
    ws = wb[M.BANK_SHEET_NAME]
    header_row = M.find_bank_header_row(ws)
    if header_row is None:
        raise ValueError("Could not find the bank's column-header row (expected קטגוריה/יצרן/תיאור).")

    row_num = ws.max_row + 1
    today = date.today().strftime("%d.%m.%Y")
    new_vals = {
        "קטגוריה": fixture["category"], "יצרן": fixture["יצרן"], "ספק": supplier,
        "סטטוס": "פעיל", "תיאור": fixture["תיאור"], "גוון": fixture.get("גוון", ""),
        "מתח_עמעום": fixture.get("מתח_עמעום", ""), "סימון": fixture.get("סימון", ""),
        "פרויקטים_מקור": fixture.get("project_name", ""), "תאריך_הוספה": today,
    }
    for c, _ in M.BANK_COLUMNS:
        if c == "תמונה":
            continue
        cell = ws.cell(row=row_num, column=M.col_index(c), value=new_vals.get(c, ""))
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    if fixture.get("image_bytes"):
        M.write_image(ws, row_num, fixture["image_bytes"], fixture["image_w"], fixture["image_h"])

    last_col_letter = openpyxl.utils.get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{row_num}"

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    save_workbook(wb, out_path)
    return row_num
