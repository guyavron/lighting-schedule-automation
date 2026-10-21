#!/usr/bin/env python3
"""
Migrate fixtures from an old project מקרא workbook into the flat בנק גופי
תאורה (fixture bank) workbook.

Reuses the same "read the actual header row, never hardcode a column letter"
principle as mikra/scripts/insert_fixture.py, because real project files
customize column order/count. Also consolidates old manual multi-row-merged
fixture entries (a whole fixture spread visually across 2-5 Excel rows) into
a single logical row — the merge is cosmetic, all the data already lives in
the top-left anchor cell of each merged column range, so no text-splicing is
needed, just reading that anchor row correctly and not re-counting its
continuation rows as separate (or empty) entries.

A row is EXCLUDED only for:
  - missing יצרן or missing תיאור (a required field wasn't filled)
  - no product photo could be matched to the row (מוצר לא זוהה חזותית)
Column order differences between files, or a תיאור that doesn't follow the
exact studio phrasing skeleton, do NOT exclude a row.

Duplicates (same יצרן + same model key extracted from תיאור) already in the
bank are merged automatically: the more complete entry (more non-empty
fields) wins, and the row's פרויקטים_מקור list accumulates every source
project.

Usage:
    python3 migrate_mikra_to_bank.py <source_mikra.xlsx> <bank.xlsx> \
        [--supplier "יצרן=ספק" ...] [--out /path/to/updated_bank.xlsx]

If a manufacturer has no --supplier mapping given, ספק is left blank for its
rows and the manufacturer is listed in the report for follow-up.
"""
import argparse
import copy
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "skill" / "scripts"))
from xlsx_io import WorkbookError, open_workbook, save_workbook  # noqa: E402
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")


SOURCE_SHEET_NAME = "אומדן"
SOURCE_HEADER_ROW = 8
CATEGORY_TEXT = {
    "תקרה": "מנורות הזנה מהתקרה",
    "קיר": "מנורות הזנה מהקיר",
    "רצפה": "מנורות הזנה מהרצפה",
}
SUMMARY_MARKER = 'סה"כ כולל'

BANK_SHEET_NAME = "בנק"
BANK_COLUMNS = [  # canonical name, header text (must match build_bank_template.py)
    ("קטגוריה", "קטגוריה"),
    ("יצרן", "יצרן"),
    ("ספק", "ספק"),
    ("סטטוס", "סטטוס"),
    ("תיאור", "תיאור"),
    ("גוון", "גוון"),
    ("מתח_עמעום", "מתח / פרוטוקול עמעום"),
    ("סימון", "סימון (מהפרויקט המקורי)"),
    ("תמונה", "תמונה"),
    ("פרויקטים_מקור", "פרויקט/ים מקור"),
    ("תאריך_הוספה", "תאריך הוספה לבנק"),
]

IMG_MAX_W_PX = 130
IMG_MAX_H_PX = 110
DEFAULT_ROW_H_PT = 20

BANK_HEADER_TEXT = {text for _, text in BANK_COLUMNS}


def normalize(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


# Fields whose LINE BREAKS are part of the studio format (one spec fragment per
# line; voltage and dimming protocol on two lines). normalize() flattens them
# into one long line -- right for comparing and for dedup keys, wrong for
# storing: a fixture exported back from the bank would no longer match the
# format of a row written straight from a PDF.
MULTILINE_FIELDS = {"תיאור", "מתח_עמעום"}


def normalize_text(v):
    """Like normalize(), but keeps line breaks (tidies spaces inside each line, drops empty lines)."""
    if v in (None, ""):
        return ""
    lines = (re.sub(r"[ \t\u00a0]+", " ", line).strip() for line in str(v).splitlines())
    return "\n".join(line for line in lines if line)


def find_bank_header_row(ws, search_rows=20):
    """The bank sheet has a branded header block above the column-header row
    (logo, title — added 23.08.2026 to match מקרא's own look), so the column
    header isn't assumed to sit at row 1 either — same "search, don't assume"
    principle as mikra/scripts/insert_fixture.py's HEADER_ROW handling.
    Returns the first row (within the first `search_rows` rows) whose cells
    contain at least the קטגוריה, יצרן and תיאור header texts, or None."""
    required = {"קטגוריה", "יצרן", "תיאור"}
    for row in range(1, search_rows + 1):
        found = {normalize(ws.cell(row=row, column=c).value) for c in range(1, ws.max_column + 1)}
        if required <= found:
            return row
    return None


def classify_header(norm):
    if norm.startswith("סעיף"):
        return "seif"
    if norm == "סימון":
        return "סימון"
    if norm == "תיאור":
        return "תיאור"
    if norm == "יצרן":
        return "יצרן"
    if norm == "גוון":
        return "גוון"
    if norm == "תמונה":
        return "תמונה"
    if "מתח" in norm or "עמעום" in norm:
        return "מתח_עמעום"
    return None


def build_header_map(ws, header_row=SOURCE_HEADER_ROW):
    header_map = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(row=header_row, column=col).value
        if raw is None:
            continue
        canon = classify_header(normalize(raw))
        if canon and canon not in header_map:
            header_map[canon] = col
    return header_map


def find_header_row(ws, search_rows=25):
    """Real project files sometimes insert extra note/instruction rows above
    the header (seen in practice: a DALI-coordination note pushed the header
    from row 8 to row 13) — so the header ROW position isn't assumed fixed
    either, only searched for. Returns (row_number, header_map) for the first
    row (within the first `search_rows` rows) that has both יצרן and תיאור
    recognized, or (None, {}) if none does."""
    for row in range(1, search_rows + 1):
        hmap = build_header_map(ws, header_row=row)
        if "יצרן" in hmap and "תיאור" in hmap:
            return row, hmap
    return None, {}


def find_category_blocks(ws, header_row=SOURCE_HEADER_ROW):
    header_rows = {}
    summary_row = None
    max_col = ws.max_column
    for row in range(header_row + 1, ws.max_row + 1):
        for col in range(1, max_col + 1):
            val = ws.cell(row=row, column=col).value
            if not val:
                continue
            val_clean = normalize(val)
            for label, text in CATEGORY_TEXT.items():
                if label not in header_rows and val_clean == text.strip():
                    header_rows[label] = row
            if summary_row is None and SUMMARY_MARKER in val_clean:
                summary_row = row
    if not header_rows:
        raise ValueError("No category headers found in source workbook.")
    if summary_row is None:
        summary_row = ws.max_row + 1
    ordered = sorted(header_rows.items(), key=lambda kv: kv[1])
    blocks = {}
    for i, (label, header_row) in enumerate(ordered):
        next_row = ordered[i + 1][1] if i + 1 < len(ordered) else summary_row
        blocks[label] = (header_row + 1, next_row - 1)
    return blocks


def project_name_from_header(ws):
    raw = ws.cell(row=2, column=3).value
    if not raw:
        return "לא ידוע"
    txt = normalize(raw)
    if ":" in txt:
        txt = txt.split(":", 1)[1].strip()
    return txt or "לא ידוע"


def images_by_anchor_row(ws):
    """0-indexed anchor row -> Image object."""
    out = {}
    for img in getattr(ws, "_images", []):
        anchor = img.anchor
        frm = getattr(anchor, "_from", None)
        if frm is None:
            continue
        out.setdefault(frm.row, img)
    return out


# Characters allowed inside a model name. Ø (diameter), ° (beam angle), × and +
# are part of real model names ("Arca Mini Ø55", "Kito 24°"): leaving them out
# cut the key short, so "Arca Mini Ø55" and "Arca Mini Ø55 IP65" both became
# "arca mini" and two different products were merged into one bank row.
MODEL_RE = re.compile(r"מדגם\s+([A-Za-z0-9֐-׿][A-Za-z0-9֐-׿ \-\./Øø°×+]{2,40})")


def extract_model_key(description):
    """Best-effort model identifier from a תיאור string, for duplicate
    detection. Not a guaranteed-unique key — matches are still a heuristic,
    flagged in the report for human sanity-check."""
    norm = normalize(description)
    m = MODEL_RE.search(norm)
    if m:
        return normalize(m.group(1)).split("\n")[0][:40].lower()
    # fallback: first ~6 words of the description
    return " ".join(norm.split()[:6]).lower()


def read_source_fixtures(source_path):
    """Returns (fixtures, excluded, project_name).
    fixtures: list of dicts with canonical fields + image bytes/size + row info.
    excluded: list of dicts describing rows skipped and why.
    """
    wb = open_workbook(source_path, data_only=True)
    if SOURCE_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"No '{SOURCE_SHEET_NAME}' sheet in this workbook.")
    ws = wb[SOURCE_SHEET_NAME]
    header_row, header_map = find_header_row(ws)
    if header_row is None:
        raise ValueError("Could not find a header row with both יצרן and תיאור in the first 25 rows.")
    blocks = find_category_blocks(ws, header_row=header_row)
    project_name = project_name_from_header(ws)
    img_map = images_by_anchor_row(ws)

    anchor_col = header_map.get("seif") or header_map["יצרן"]

    fixtures, excluded = [], []
    for category, (first, last) in blocks.items():
        row = first
        while row <= last:
            cell = ws.cell(row=row, column=anchor_col)
            if isinstance(cell, MergedCell) or cell.value in (None, ""):
                row += 1
                continue
            # this is a fixture anchor row (possibly a merged multi-row block)
            def val(field):
                col = header_map.get(field)
                if not col:
                    return ""
                v = ws.cell(row=row, column=col).value
                return normalize_text(v) if field in MULTILINE_FIELDS else normalize(v)

            yatzran = val("יצרן")
            tiur = val("תיאור")
            img = img_map.get(row - 1)  # 0-indexed

            info = {
                "row": row, "category": category,
                "יצרן": yatzran, "תיאור": tiur,
                "גוון": val("גוון"), "מתח_עמעום": val("מתח_עמעום"),
                "סימון": val("סימון"),
            }

            reasons = []
            if not yatzran:
                reasons.append("חסר יצרן")
            if not tiur:
                reasons.append("חסר תיאור")
            if img is None:
                reasons.append("לא נמצאה תמונה מתאימה (מוצר לא זוהה חזותית)")

            if reasons:
                excluded.append({**info, "reasons": reasons})
            else:
                info["image_bytes"] = img._data()
                info["image_w"] = img.width
                info["image_h"] = img.height
                fixtures.append(info)

            # advance past this fixture's merged block: find how many rows
            # its anchor column spans via the sheet's merged ranges
            span = 1
            for mc in ws.merged_cells.ranges:
                if mc.min_col <= anchor_col <= mc.max_col and mc.min_row == row:
                    span = max(span, mc.max_row - mc.min_row + 1)
            row += span

    return fixtures, excluded, project_name


def scaled_size(w, h, max_w=IMG_MAX_W_PX, max_h=IMG_MAX_H_PX):
    ratio = min(max_w / w, max_h / h)
    return max(1, round(w * ratio)), max(1, round(h * ratio))


def col_letter(canon):
    for i, (c, _) in enumerate(BANK_COLUMNS, start=1):
        if c == canon:
            return get_column_letter(i)
    raise KeyError(canon)


def col_index(canon):
    for i, (c, _) in enumerate(BANK_COLUMNS, start=1):
        if c == canon:
            return i
    raise KeyError(canon)


def read_bank_rows(ws, header_row):
    """Returns list of dicts (row_num included) for existing bank data rows."""
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = {c: ws.cell(row=r, column=col_index(c)).value for c, _ in BANK_COLUMNS}
        if not any(v not in (None, "") for v in vals.values()):
            continue
        vals["_row"] = r
        rows.append(vals)
    return rows


def completeness(vals):
    fields = ["ספק", "סטטוס", "גוון", "מתח_עמעום", "סימון"]
    return sum(1 for f in fields if normalize(vals.get(f, "")))


def resize_image_bytes(image_bytes, target_w, target_h):
    """Actually re-encode pixels at the display size, not just set XLImage's
    width/height attributes — those only scale how Excel RENDERS the image,
    the full-resolution original bytes still get embedded either way. A bank
    with 100+ photos at full manufacturer-PDF resolution (often 500-900px,
    tens/hundreds of KB each) balloons the file into the tens of MB.

    Always flattens to RGB (white background) and re-encodes as JPEG: this is
    a browsing/filtering thumbnail, not an archival copy, and lossless PNG
    re-encoding of a resized photo is both slower and — for photographic
    content — often BIGGER than the tiny original (smoothing introduces many
    intermediate colours a flat-colour PNG didn't have to compress)."""
    from PIL import Image
    import io
    src = Image.open(io.BytesIO(image_bytes))
    if src.mode in ("RGBA", "LA", "P"):
        src = src.convert("RGBA")
        bg = Image.new("RGB", src.size, (255, 255, 255))
        bg.paste(src, mask=src.split()[-1])
        src = bg
    else:
        src = src.convert("RGB")
    resized = src.resize((max(1, target_w), max(1, target_h)), Image.LANCZOS)
    out = io.BytesIO()
    resized.save(out, format="JPEG", quality=78)
    return out.getvalue()


def write_image(ws, row, image_bytes, orig_w, orig_h):
    import io
    w, h = scaled_size(orig_w, orig_h)
    small_bytes = resize_image_bytes(image_bytes, w, h)
    img = XLImage(io.BytesIO(small_bytes))
    img.width, img.height = w, h
    anchor = f"{col_letter('תמונה')}{row}"
    ws.add_image(img, anchor)
    row_pt = max(DEFAULT_ROW_H_PT, round(h * 0.8) + 6)
    ws.row_dimensions[row].height = row_pt


def remove_image_at_row(ws, row):
    target_col0 = col_index("תמונה") - 1
    keep = []
    for img in getattr(ws, "_images", []):
        frm = getattr(img.anchor, "_from", None)
        if frm is not None and frm.row == row - 1 and frm.col == target_col0:
            continue
        keep.append(img)
    ws._images = keep


def migrate(source_path, bank_path, out_path, supplier_map):
    fixtures, excluded, project_name = read_source_fixtures(source_path)

    wb = open_workbook(bank_path, sheet=BANK_SHEET_NAME)
    ws = wb[BANK_SHEET_NAME]
    header_row = find_bank_header_row(ws)
    if header_row is None:
        raise ValueError("Could not find the bank's column-header row (expected קטגוריה/יצרן/תיאור).")
    existing = read_bank_rows(ws, header_row)
    existing_index = {}
    for e in existing:
        key = (normalize(e.get("יצרן", "")).lower(), extract_model_key(e.get("תיאור", "")))
        existing_index[key] = e

    added, merged, needs_supplier = [], [], set()
    today = date.today().strftime("%d.%m.%Y")

    next_row = (ws.max_row + 1) if existing else header_row + 1

    for fx in fixtures:
        yatzran_norm = normalize(fx["יצרן"])
        supplier = supplier_map.get(yatzran_norm.lower(), "")
        if not supplier:
            needs_supplier.add(fx["יצרן"])
        key = (yatzran_norm.lower(), extract_model_key(fx["תיאור"]))
        new_vals = {
            "קטגוריה": fx["category"], "יצרן": fx["יצרן"], "ספק": supplier,
            "סטטוס": "פעיל", "תיאור": fx["תיאור"], "גוון": fx["גוון"],
            "מתח_עמעום": fx["מתח_עמעום"], "סימון": fx["סימון"],
            "פרויקטים_מקור": project_name, "תאריך_הוספה": today,
        }

        if key in existing_index:
            old = existing_index[key]
            old_score = completeness(old)
            new_score = completeness(new_vals)
            row_num = old["_row"]
            projects = set(p.strip() for p in normalize(old.get("פרויקטים_מקור", "")).split(",") if p.strip())
            projects.add(project_name)
            if new_score > old_score:
                for c, _ in BANK_COLUMNS:
                    if c in ("פרויקטים_מקור", "תאריך_הוספה"):
                        continue
                    ws.cell(row=row_num, column=col_index(c), value=new_vals.get(c, ""))
                remove_image_at_row(ws, row_num)
                write_image(ws, row_num, fx["image_bytes"], fx["image_w"], fx["image_h"])
            ws.cell(row=row_num, column=col_index("פרויקטים_מקור"), value=", ".join(sorted(projects)))
            merged.append({**fx, "bank_row": row_num, "kept": "new" if new_score > old_score else "existing"})
            continue

        row_num = next_row
        for c, _ in BANK_COLUMNS:
            if c == "תמונה":
                continue
            cell = ws.cell(row=row_num, column=col_index(c), value=new_vals.get(c, ""))
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        write_image(ws, row_num, fx["image_bytes"], fx["image_w"], fx["image_h"])
        existing_index[key] = {**new_vals, "_row": row_num}
        added.append({**fx, "bank_row": row_num})
        next_row += 1

    last_row = next_row - 1
    # use the sheet's actual column count, not len(BANK_COLUMNS) -- a column
    # appended outside this list (e.g. export_selection.py's "בחירה לפרויקט"
    # marking column) must stay inside the filter range too.
    last_col_letter = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{max(last_row, header_row)}"

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    save_workbook(wb, out_path)

    return {
        "project_name": project_name,
        "added": added, "merged": merged, "excluded": excluded,
        "needs_supplier": sorted(needs_supplier),
    }


def parse_supplier_args(pairs):
    out = {}
    for p in pairs or []:
        if "=" not in p:
            continue
        k, v = p.split("=", 1)
        out[normalize(k).lower()] = normalize(v)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source_xlsx")
    ap.add_argument("bank_xlsx")
    ap.add_argument("--supplier", action="append", default=[], help='"יצרן=ספק", repeatable')
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    out_path = args.out or args.bank_xlsx
    supplier_map = parse_supplier_args(args.supplier)
    report = migrate(args.source_xlsx, args.bank_xlsx, out_path, supplier_map)

    print(f"פרויקט מקור: {report['project_name']}")
    print(f"נוספו {len(report['added'])} שורות חדשות:")
    for f in report["added"]:
        print(f"  + [{f['category']}] {f['יצרן']} — {f['תיאור'][:50].replace(chr(10), ' · ')}")
    print(f"אוחדו {len(report['merged'])} כפילויות:")
    for f in report["merged"]:
        print(f"  = [{f['category']}] {f['יצרן']} — {f['תיאור'][:50].replace(chr(10), ' · ')} (bank row {f['bank_row']}, kept={f['kept']})")
    print(f"נפסלו {len(report['excluded'])} שורות:")
    for f in report["excluded"]:
        print(f"  - row {f['row']} [{f['category']}] {f.get('יצרן','')} — {f.get('תיאור','')[:40].replace(chr(10), ' · ')!r}: {', '.join(f['reasons'])}")
    if report["needs_supplier"]:
        print("יצרנים ללא ספק ידוע (ספק נשאר ריק):")
        for y in report["needs_supplier"]:
            print(f"  ? {y}")
    print(f"\nנשמר: {out_path}")


if __name__ == "__main__":
    try:
        main()
    except (WorkbookError, ValueError) as e:
        raise SystemExit(f"שגיאה: {e}")
