#!/usr/bin/env python3
"""
Build the fixture selection page from the bank workbook. Re-run whenever the
bank grows (bank/rebuild_pages.py does this automatically).

The page is a snapshot: every bank row and a downscaled thumbnail are embedded
in the HTML itself, so it opens in any browser with no server, account or
connector. The page only SELECTS fixtures -- it produces the row list that
bank/export_selection.py --rows turns into a real legend workbook. (A web page
can't write the legend's merged cells, colour bars, row heights and anchored
images; that stays in Python.)

    python pages/build_pages.py --bank examples/out/bank.xlsx --outdir docs --out-name index.html --demo-note
"""
import argparse, base64, io, json, os, sys
from pathlib import Path

import openpyxl
from PIL import Image

# עברית בפלט גם בווינדוס, כשהפלט נקרא דרך pipe (ברירת המחדל שם היא cp1252)
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")


HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
sys.path.insert(0, str(PROJECT / "bank"))
import migrate_mikra_to_bank as M  # noqa: E402  (bank column layout + header-row search)

THUMB_PX = 150          # max thumbnail size on the page (the bank itself keeps up to 130x110)
JPEG_QUALITY = 72
# page key -> canonical bank column. Looked up by header text, never by a fixed
# column number: a column added to the bank would otherwise silently shift every
# field on the page by one.
FIELDS = {"cat": "קטגוריה", "man": "יצרן", "sup": "ספק", "st": "סטטוס", "desc": "תיאור",
          "fin": "גוון", "volt": "מתח_עמעום", "mark": "סימון", "proj": "פרויקטים_מקור"}
TEMPLATE = "templates/picker.html"
OUT_NAME = "picker.html"

STANDALONE_HEAD = """<!doctype html>
<html lang="he" dir="rtl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
"""
STANDALONE_TAIL = "\n</html>\n"


def cell(ws, r, c):
    v = ws.cell(row=r, column=c).value
    return "" if v is None else str(v).strip()


def thumbnails(ws):
    """row number -> data URI of a downscaled thumbnail."""
    out = {}
    for im in ws._images:
        row = im.anchor._from.row + 1
        try:
            raw = im.ref.getvalue() if hasattr(im.ref, "getvalue") else open(im.ref, "rb").read()
            pic = Image.open(io.BytesIO(raw))
            try:
                pic.draft("RGB", (THUMB_PX, THUMB_PX))   # fast JPEG decode
            except Exception:
                pass
            pic = pic.convert("RGB")
            pic.thumbnail((THUMB_PX, THUMB_PX))
            buf = io.BytesIO()
            pic.save(buf, "JPEG", quality=JPEG_QUALITY, optimize=True)
            out[row] = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
        except Exception as e:
            print(f"  אזהרה: תמונה בשורה {row} לא נקראה ({e})", file=sys.stderr)
    return out


def read_bank(bank_path):
    wb = openpyxl.load_workbook(bank_path)
    if M.BANK_SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"בקובץ {bank_path} אין גיליון בשם '{M.BANK_SHEET_NAME}' — זה לא קובץ בנק.")
    ws = wb[M.BANK_SHEET_NAME]
    header_row = M.find_bank_header_row(ws)
    if header_row is None:
        raise SystemExit("לא נמצאה שורת הכותרות של הבנק (קטגוריה / יצרן / תיאור).")
    cols = {}
    for c in range(1, ws.max_column + 1):
        text = M.normalize(ws.cell(row=header_row, column=c).value)
        for canon, header_text in M.BANK_COLUMNS:
            if text == header_text and canon not in cols:
                cols[canon] = c
    missing = [h for h in ("קטגוריה", "יצרן", "תיאור") if h not in cols]
    if missing:
        raise SystemExit(f"חסרות עמודות חובה בבנק: {missing}")
    imgs = thumbnails(ws)
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        rec = {k: (cell(ws, r, cols[canon]) if canon in cols else "") for k, canon in FIELDS.items()}
        if not (rec["man"] or rec["desc"]):
            continue
        rec["r"] = r
        rec["img"] = imgs.get(r, "")
        rows.append(rec)
    wb.close()
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bank", default=str(PROJECT / "bank.xlsx"))
    ap.add_argument("--outdir", default=str(PROJECT))
    ap.add_argument("--out-name", default=OUT_NAME)
    ap.add_argument("--fragment", action="store_true",
                    help="write only the page body (for publishing as a Claude artifact, which adds "
                         "its own <html>/<head>); default is a full HTML document that opens in any browser")
    ap.add_argument("--demo-note", action="store_true",
                    help="show the 'fictional demo data' banner and the equivalent CLI command")
    args = ap.parse_args()

    rows = read_bank(args.bank)
    with_img = sum(1 for x in rows if x["img"])
    print(f"נקראו {len(rows)} גופים מהבנק ({with_img} עם תמונה).")
    # "</" is escaped so a description can never close the <script> block early
    payload = json.dumps(rows, ensure_ascii=False).replace("</", "<\\/")

    tmpl = (HERE / TEMPLATE).read_text(encoding="utf8")
    if "__DATA__" not in tmpl:
        raise SystemExit(f"{TEMPLATE} חסר הסימון __DATA__.")
    html = tmpl.replace("__DATA__", payload).replace("__DEMO__", "true" if args.demo_note else "false")
    if not args.fragment:
        html = STANDALONE_HEAD + html.replace("</style>", "</style>\n</head>\n<body>", 1) \
            + "\n</body>" + STANDALONE_TAIL
    out_path = Path(args.outdir) / args.out_name
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf8")
    print(f"נכתב: {out_path}")


if __name__ == "__main__":
    main()
